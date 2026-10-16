#!/usr/bin/env python3
"""核对排版结果：人数、体检编号、姓名、正文、表头、页数、避头尾、文字是否越界。

用法：python3 verify_report.py 原始文件.xlsx 输出.xlsx 输出.pdf layout_plan.json
"""
import json
import os
import re
import sys
from collections import defaultdict

import openpyxl
import pdfplumber

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_report import (COLS, FIRST_DATA_ROW, HEADER_ROW, LINE_HEAD_FORBIDDEN,  # noqa: E402
                          read_source)

HEADERS = [h for _, h, _, _, _ in COLS]


def squash(s):
    return re.sub(r"[\s\u2060]+", "", str(s))


def main(src, xlsx, pdf_path, plan_path):
    errors = []
    plan = json.load(open(plan_path, encoding="utf-8"))
    annex, title, info, headers, people = read_source(src)

    # ---- 1. Excel 与原表逐人核对 ------------------------------------------
    raw_ws = openpyxl.load_workbook(src).active
    src_rows = [r for r in range(FIRST_DATA_ROW, raw_ws.max_row + 1)
                if any(raw_ws.cell(r, c).value not in (None, "") for c in range(1, 11))]
    ws = openpyxl.load_workbook(xlsx).active
    out_rows = [r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")]
    if len(src_rows) != len(out_rows):
        errors.append(f"人数不一致：原表 {len(src_rows)}，输出 {len(out_rows)}")
    for i, (sr, orow) in enumerate(zip(src_rows, out_rows)):
        for c in range(1, 11):
            a, b = raw_ws.cell(sr, c).value, ws.cell(orow, c).value
            if c in (8, 10):
                # 正文只允许：删除空白/空行、用“。”连接各检查项目、序号统一为“1.”
                def core(t):
                    return re.sub(r"[\s。.]", "", str(t))
                if core(a) != core(b):
                    errors.append(f"第 {i + 1} 人 第 {c} 列正文与原表不一致")
                if "。。" in str(b) or str(b).count(".") < str(a).count("."):
                    errors.append(f"第 {i + 1} 人 第 {c} 列标点异常")
            else:
                if squash(a) != squash(b):
                    errors.append(f"第 {i + 1} 人 第 {c} 列：{a!r} → {b!r}")
    for c, h in enumerate(HEADERS, start=1):
        if ws.cell(HEADER_ROW, c).value != h:
            errors.append(f"表头第 {c} 列为 {ws.cell(HEADER_ROW, c).value!r}")
        if not ws.cell(HEADER_ROW, c).font.b or ws.cell(HEADER_ROW, c).font.sz != 10:
            errors.append(f"表头第 {c} 列字体不是 10 号加粗")
    if ws.print_title_rows != f"${HEADER_ROW}:${HEADER_ROW}":
        errors.append(f"打印标题行为 {ws.print_title_rows}")
    breaks = [b.id for b in ws.row_breaks.brk]
    exp_breaks = []
    pos = FIRST_DATA_ROW - 1
    for pg in plan["pages"][:-1]:
        pos += len(pg)
        exp_breaks.append(pos)
    if breaks != exp_breaks:
        errors.append("Excel 分页符与排版计划不一致")
    over = [r for r in out_rows if ws.row_dimensions[r].height > 409]
    if over:
        errors.append(f"行高超过 Excel 上限 409：{over}")

    # ---- 2. PDF 逐页核对 ------------------------------------------------
    pdf = pdfplumber.open(pdf_path)
    if len(pdf.pages) != len(plan["pages"]):
        errors.append(f"PDF 页数 {len(pdf.pages)} ≠ 计划 {len(plan['pages'])}")
    id_to_person = {p["vals"][0]: p for p in people}
    head_issues = []
    seen_ids = []
    name_ok = 0
    for pno, (page, ids) in enumerate(zip(pdf.pages, plan["pages"]), start=1):
        chars = [c for c in page.chars if c["text"].strip() and c["text"] != "\u2060"]
        vedges = sorted({round(e["x0"], 1) for e in page.edges if e["orientation"] == "v"})
        hedges = [e for e in page.edges if e["orientation"] == "h"]
        if len(vedges) < 11:
            errors.append(f"第 {pno} 页竖线数量异常：{len(vedges)}")
            continue
        col_x = vedges[-11:]                     # 表格 11 条竖线
        text_all = squash("".join(c["text"] for c in sorted(chars, key=lambda c: (round(c["top"]), c["x0"]))))
        # 首页信息 / 后续页不重复
        first_page_bits = [annex, title] + [x for pair in info for x in pair]
        for bit in first_page_bits:
            present = squash(bit) in text_all
            if pno == 1 and not present:
                errors.append(f"首页缺少：{bit}")
            if pno > 1 and squash(bit) in text_all and bit != annex:
                errors.append(f"第 {pno} 页重复了：{bit}")
        # 表头所在行：含“体检编号”的那一行
        hdr = [c for c in chars if c["text"] == "体" and col_x[0] < c["x0"] < col_x[1]]
        if not hdr:
            errors.append(f"第 {pno} 页没有表头")
            continue
        hdr_mid = max(c["top"] for c in hdr)          # 表头“体检编号”的“体”（首页信息栏在其上方）
        # 贯穿整表的横线：表头上下边及各人行的底边
        all_y = sorted({round(e["top"], 1) for e in hedges
                        if e["x0"] <= col_x[0] + 1 and e["x1"] >= col_x[-1] - 1})
        hdr_top = max(y for y in all_y if y <= hdr_mid)
        lines_y = [y for y in all_y if y > hdr_mid]
        hdr_bottom = lines_y[0]
        header_text = []
        for ci in range(10):
            cs = [c for c in chars if col_x[ci] < c["x0"] < col_x[ci + 1] and hdr_top < c["top"] < hdr_bottom]
            cs.sort(key=lambda c: (round(c["top"]), c["x0"]))
            header_text.append("".join(c["text"] for c in cs))
            fs = {round(c["size"], 1) for c in cs}
            if fs != {10.0}:
                head_issues.append((pno, ci, fs))
        if header_text != [squash(h) for h in HEADERS]:
            errors.append(f"第 {pno} 页表头不完整：{header_text}")
        # 数据行
        rows_y = lines_y  # 表头底边、各人行底边
        if len(rows_y) != len(ids) + 1:
            errors.append(f"第 {pno} 页行线数量 {len(rows_y)} 与人数 {len(ids)} 不符")
            continue
        for k, pid in enumerate(ids):
            y0, y1 = rows_y[k], rows_y[k + 1]
            person = id_to_person[pid]
            seen_ids.append(pid)
            for ci in range(10):
                x0, x1 = col_x[ci], col_x[ci + 1]
                cs = [c for c in chars if x0 - 1 < (c["x0"] + c["x1"]) / 2 < x1 + 1 and y0 - 30 < c["top"] < y1 + 30]
                inside = [c for c in cs if c["top"] >= y0 - 0.5 and c["bottom"] <= y1 + 0.5
                          and c["x0"] >= x0 - 0.5 and c["x1"] <= x1 + 0.5]
                stray = [c for c in cs if c not in inside and y0 <= (c["top"] + c["bottom"]) / 2 <= y1]
                if stray:
                    errors.append(f"第 {pno} 页 {pid} 第 {ci + 1} 列文字越出单元格")
                # 按行分组
                lines = defaultdict(list)
                for c in inside:
                    lines[round(c["top"] * 2) / 2].append(c)
                ordered = [sorted(v, key=lambda c: c["x0"]) for _, v in sorted(lines.items())]
                got = "".join(c["text"] for ln in ordered for c in ln)
                want = squash(person["vals"][ci])
                if squash(got) != want:
                    errors.append(f"第 {pno} 页 {pid} 第 {ci + 1} 列文字不完整："
                                  f"PDF {len(squash(got))} 字 / 应 {len(want)} 字")
                elif ci == 1:
                    name_ok += 1
                for ln in ordered:
                    if ln[0]["text"] in LINE_HEAD_FORBIDDEN:
                        ctx = "".join(c["text"] for c in ln[:12])
                        errors.append(f"第 {pno} 页 {pid} 第 {ci + 1} 列行首标点：{ctx}")
                # 数字或英文单词被硬拆到两行（原文此处没有空格/换行）
                full = str(person["vals"][ci])
                j = 0
                for li, ln in enumerate(ordered):
                    for ki, c in enumerate(ln):
                        gap = ""
                        while j < len(full) and full[j] != c["text"] and full[j] in " \n\u2060":
                            gap += full[j]
                            j += 1
                        if j >= len(full) or full[j] != c["text"]:
                            break
                        if ki == 0 and li > 0 and not gap.replace("\u2060", ""):
                            a, b = ordered[li - 1][-1]["text"], c["text"]
                            if (re.match(r"[0-9.]", a) and re.match(r"[0-9.]", b)) or \
                                    (re.match(r"[A-Za-z]", a) and re.match(r"[A-Za-z]", b)):
                                ctx = "".join(x["text"] for x in ordered[li - 1][-6:]) + "|" + \
                                    "".join(x["text"] for x in ln[:6])
                                errors.append(f"第 {pno} 页 {pid} 第 {ci + 1} 列数字/单词被拆行：{ctx}")
                        j += 1
    if seen_ids != [p["vals"][0] for p in people]:
        errors.append("PDF 中人员顺序/数量与原表不一致")
    if head_issues:
        errors.append(f"表头字号异常：{head_issues[:5]}")

    print(f"原表人数 {len(src_rows)}，Excel 人数 {len(out_rows)}，PDF 人数 {len(seen_ids)}"
          f"（体检编号逐页核对，姓名吻合 {name_ok} 人），"
          f"PDF 页数 {len(pdf.pages)}（计划 {len(plan['pages'])}，单人页 "
          f"{sum(len(p) == 1 for p in plan['pages'])}，双人页 {sum(len(p) == 2 for p in plan['pages'])}）")
    if errors:
        print(f"发现 {len(errors)} 个问题：")
        for e in errors[:80]:
            print("  -", e)
        sys.exit(1)
    print("全部核对通过。")


if __name__ == "__main__":
    main(*sys.argv[1:5])
