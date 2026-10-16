#!/usr/bin/env python3
"""职业健康检查一览表（附件8）排版脚本。

读取原始一览表 xlsx，整理文字并按打印要求重新排版，输出：
  * 可编辑的 Excel（固定行高、手动分页、每页重复表头）
  * 由同一 Excel 经 LibreOffice 导出的 PDF（最终打印文件）

用法：
  python3 build_report.py 原始文件.xlsx 输出目录 [输出文件名（不含扩展名）]

依赖：openpyxl、LibreOffice（soffice 及其 Python UNO 模块）、
PDF 渲染用的“宋体”替代字体（见 make_song_font.py 与 fonts.conf）。
"""
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.page import PageMargins

FONT = "宋体"
# 工作簿“常规”样式字体。宋体 12 的数字宽度恰为 6 pt（8 px），Excel 与
# LibreOffice 换算列宽的结果完全一致，保证两者的列宽、换行与分页相同。
NORMAL_FONT = Font(name=FONT, size=12)

# ---- 版面（单位：Excel 像素，96 dpi；1 px = 0.75 pt） -----------------------
# A4 横向，左右边距 0.5 cm → 可打印宽度 287 mm ≈ 1084 px，表格总宽 1060 px 留余量。
COLS = [
    # 列, 表头, 宽度px, 正文字号, 水平对齐
    ("A", "体检编号", 66, 9, "center"),
    ("B", "姓名", 42, 9, "center"),
    ("C", "性别", 32, 9, "center"),
    ("D", "年龄", 32, 9, "center"),
    ("E", "接害\n工龄", 54, 9, "center"),
    ("F", "职业病\n危害因素", 66, 9, "center"),      # “一氧化碳、/噪声”两行
    ("G", "体检种类", 60, 9, "center"),
    ("H", "体检结果", 412, 8, "left"),
    ("I", "体检结论", 60, 9, "center"),
    ("J", "临床建议及处理意见", 236, 8, "left"),   # I+J 容纳 12 号日期
]
LONG_COLS = ("H", "J")

PAGE_H_PT = (210 - 10 - 10) / 25.4 * 72      # 上下边距各 1 cm
SAFETY_PT = 10                                 # 每页预留，防止 Excel 自动分页
MAX_ROW_PT = 409                               # Excel 单行行高上限
ROW_PAD_PT = 4                                 # 每行在 LibreOffice 最适行高上的余量
MIN_LONG_FONT = 7                              # 长文本最小字号

TITLE_ROWS = {1: 18, 2: 26, 3: 21, 4: 21}      # 首页标题区行高（pt）
HEADER_ROW = 5
HEADER_H = 28
FIRST_DATA_ROW = 6

THIN = Side(style="thin", color="000000")
ALL_BORDERS = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------- 文字整理 --
CJK_PUNCT = "，。：；、？！（）【】“”"
END_PUNCT = "。；，、：？！;:,…"
# 不得出现在行首 / 行尾的字符（中文避头尾规则，另加 %、↑↓、～ 等紧随数值的符号）
LINE_HEAD_FORBIDDEN = "，。、；：？！）】”’》」』〉〕…—～·,.;:!?)]}%％‰℃°↑↓"
LINE_TAIL_FORBIDDEN = "（【“‘《「『〈〔([{"


def norm_spaces(s):
    """去掉首尾及多余空格；中文标点两侧的空格一律删除。"""
    s = s.replace("　", " ").replace("\xa0", " ").replace("\t", " ")
    s = re.sub(r" {2,}", " ", s).strip()
    s = re.sub(r" *([%s]) *" % CJK_PUNCT, r"\1", s)
    s = re.sub(r"(?<=\d) +(?=%)", "", s)          # “40.6 %” → “40.6%”
    s = re.sub(r"。{2,}", "。", s)                 # 原文重复的句号
    return s


def norm_number(s):
    """统一条目序号：“1. 文字” / “3 文字” → “1.文字” / “3.文字”。"""
    s = re.sub(r"^(\d+)\.\s*", r"\1.", s)
    s = re.sub(r"^(\d+)\s+(?=[一-鿿])", r"\1.", s)
    return s


def join_parts(parts):
    """同一条目内的多行合并为一段，行与行之间用句号分隔。"""
    out = parts[0]
    for p in parts[1:]:
        if out[-1] in END_PUNCT or p[0] in LINE_HEAD_FORBIDDEN:
            out += p
        else:
            out += "。" + p
    return out


def split_items(text):
    lines = [norm_number(norm_spaces(l)) for l in str(text or "").replace("\r", "\n").split("\n")]
    lines = [l for l in lines if l]
    head, items = [], []
    for l in lines:
        if re.match(r"^\d+\.", l):
            items.append([l])
        elif items:
            items[-1].append(l)
        else:
            head.append(l)
    return head, items


def clean_text(text):
    """体检结果 / 临床建议：不换行，各检查项目及 1./2./3. 各小项之间用句号隔开。"""
    head, items = split_items(text)
    body = join_parts([join_parts(it) for it in items]) if items else ""
    # “【阳性发现】”等小标题直接与第 1 项接排
    return (join_parts(head) if head else "") + body


def clean_short(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return norm_spaces(str(v))


# ---------------------------------------------------------------- 读取源表 --
def read_source(path):
    ws = openpyxl.load_workbook(path).active
    annex = norm_spaces(str(ws["A1"].value or ""))
    title = norm_spaces(str(ws["A2"].value or ""))
    info = []
    for addr in ("A3", "A4"):
        parts = [p.strip() for p in re.split(r"\s{2,}", str(ws[addr].value or "").strip()) if p.strip()]
        parts = [re.sub(r"\s*-\s*", "-", norm_spaces(p)) for p in parts]
        assert len(parts) == 2, (addr, parts)
        info.append(parts)
    headers = [clean_short(ws.cell(HEADER_ROW, c).value) for c in range(1, 11)]
    people = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 11)]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        rec = [clean_short(v) for v in vals]
        rec[7] = clean_text(vals[7])
        rec[9] = clean_text(vals[9])
        people.append({"src_row": r, "vals": rec})
    return annex, title, info, headers, people


# ---------------------------------------------------------------- 写工作簿 --
def px_to_width(px):
    # Excel 列宽以“常规”字体的数字宽度为单位：宋体 12 → 8 px
    return int(px / 8 * 256) / 256


HALF_WIDTH = re.compile(r"[\x20-\x7e]")
NO_BREAK_BEFORE = set(LINE_HEAD_FORBIDDEN)
NO_BREAK_AFTER = set(LINE_TAIL_FORBIDDEN)


def estimate_lines(text, avail_px, cjk_px, half_px):
    """按宋体字宽和中文避头尾规则粗算换行行数（用于估算 Excel 屏幕显示所需行高）。"""
    total = 0
    for para in text.split("\n"):
        units = []
        for tok in re.findall(r"[\x21-\x7e]+| |.", para):
            w = sum(half_px if HALF_WIDTH.match(ch) else cjk_px for ch in tok)
            if units and (tok[0] in NO_BREAK_BEFORE or units[-1][0][-1] in NO_BREAK_AFTER):
                units[-1] = (units[-1][0] + tok, units[-1][1] + w)
            else:
                units.append((tok, w))
        lines, cur = 1, 0
        for tok, w in units:
            if cur + w <= avail_px:
                cur += w
            elif tok == " ":
                lines, cur = lines + 1, 0
            elif w > avail_px:
                for ch in tok:
                    cw = half_px if HALF_WIDTH.match(ch) else cjk_px
                    if cur + cw > avail_px:
                        lines, cur = lines + 1, 0
                    cur += cw
            else:
                lines, cur = lines + 1, w
        total += lines
    return total


def excel_height_estimate(vals, long_font):
    """保守估算 Excel 屏幕（96 dpi，字号取整到像素）显示整行所需高度（pt）。"""
    need = 0
    for c, (col, _, px, size, _) in enumerate(COLS):
        size = long_font if col in LONG_COLS else size
        em = round(size * 96 / 72)                 # 屏幕字号像素
        pitch = em + 3                             # 与 Excel 宋体默认行高一致（9 号 15 px、12 号 19 px）
        n = estimate_lines(str(vals[c]), px - 6, em, (em + 1) // 2)
        need = max(need, n * pitch * 0.75 + 4)
    return need


def round_px(pt):
    # Excel 行高按像素（0.75 pt）取整，向上取整避免被截断
    return math.ceil(pt / 0.75 - 1e-9) * 0.75


def build_workbook(annex, title, info, people, heights=None, fonts=None, breaks=None):
    wb = openpyxl.Workbook()
    wb._fonts = type(wb._fonts)([NORMAL_FONT])
    wb._named_styles["Normal"].font = NORMAL_FONT
    ws = wb.active
    ws.title = "职业健康检查一览表"
    for col, _, px, _, _ in COLS:
        ws.column_dimensions[col].width = px_to_width(px)

    # 首页：附件名称、标题、受检单位/体检单位及日期
    ws["A1"] = annex
    ws["A1"].font = Font(name=FONT, size=12)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A2:J2")
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, size=16, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    for i, (left, right) in enumerate(info):
        r = 3 + i
        ws.merge_cells(f"A{r}:H{r}")
        ws.merge_cells(f"I{r}:J{r}")
        ws[f"A{r}"] = left
        ws[f"I{r}"] = right
        for addr in (f"A{r}", f"I{r}"):
            ws[addr].font = Font(name=FONT, size=12)
            ws[addr].alignment = Alignment(horizontal="left", vertical="center", indent=0)
    # 两行信息：完整外框 + 中间横线，日期左侧不加竖线
    for r in (3, 4):
        for c in range(1, 11):
            ws.cell(r, c).border = Border(
                left=THIN if c == 1 else None,
                right=THIN if c == 10 else None,
                top=THIN,
                bottom=THIN,
            )
    for r, h in TITLE_ROWS.items():
        ws.row_dimensions[r].height = h

    # 表头
    for c, (_, head, _, _, _) in enumerate(COLS, start=1):
        cell = ws.cell(HEADER_ROW, c, head)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = ALL_BORDERS
    ws.row_dimensions[HEADER_ROW].height = HEADER_H

    # 正文
    for i, p in enumerate(people):
        r = FIRST_DATA_ROW + i
        for c, (col, _, _, size, halign) in enumerate(COLS, start=1):
            v = p["vals"][c - 1]
            if col == "D" and v.isdigit():
                v = int(v)
            cell = ws.cell(r, c, v)
            if col in LONG_COLS and fonts:
                size = fonts[i]
            cell.font = Font(name=FONT, size=size)
            cell.alignment = Alignment(horizontal=halign, vertical="center", wrap_text=True)
            cell.border = ALL_BORDERS
            if col == "A":
                cell.number_format = "@"
        if heights:
            ws.row_dimensions[r].height = heights[i]

    last = FIRST_DATA_ROW + len(people) - 1
    # 打印设置
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.print_area = f"A1:J{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.scale = 100
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_margins = PageMargins(left=0.5 / 2.54, right=0.5 / 2.54, top=1 / 2.54,
                                  bottom=1 / 2.54, header=0, footer=0)
    ws.print_options.horizontalCentered = True
    ws.oddHeader.left.text = ws.oddHeader.center.text = ws.oddHeader.right.text = None
    ws.oddFooter.left.text = ws.oddFooter.center.text = ws.oddFooter.right.text = None
    if breaks:
        ws.row_breaks = RowBreak()
        for row in breaks:
            ws.row_breaks.append(Break(id=row))
    return wb


# ------------------------------------------------------- LibreOffice 测量/导出 --
class Office:
    """通过 UNO 驱动无界面 LibreOffice。

    LibreOffice 导入 xlsx 时会关闭单元格的“对行首行尾应用禁则”，导致逗号、
    句号、“】”等出现在行首；这里统一重新打开，再测量行高或导出 PDF。
    """

    def __enter__(self):
        import socket
        import time

        import uno
        self.uno = uno
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        self.profile = tempfile.mkdtemp(prefix="lo-profile-")
        conn = f"socket,host=127.0.0.1,port={port};urp;"
        self.proc = subprocess.Popen(
            ["soffice", "--headless", "--invisible", "--norestore", "--nologo",
             f"-env:UserInstallation=file://{self.profile}", f"--accept={conn}"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(240):
            try:
                ctx = resolver.resolve(f"uno:{conn}StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.25)
        else:
            raise RuntimeError("LibreOffice 未能启动")
        self.desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        return self

    def __exit__(self, *exc):
        try:
            self.desktop.terminate()
        except Exception:
            pass
        try:
            self.proc.wait(timeout=60)
        except subprocess.TimeoutExpired:
            self.proc.kill()
        shutil.rmtree(self.profile, ignore_errors=True)

    def _prop(self, name, value):
        p = self.uno.createUnoStruct("com.sun.star.beans.PropertyValue")
        p.Name, p.Value = name, value
        return p

    def open(self, path, n_rows):
        url = self.uno.systemPathToFileUrl(str(Path(path).resolve()))
        doc = self.desktop.loadComponentFromURL(url, "_blank", 0, (self._prop("Hidden", True),))
        sheet = doc.Sheets.getByIndex(0)
        sheet.getCellRangeByName(f"A1:J{n_rows}").setPropertyValue("ParaIsForbiddenRules", True)
        loc = self.uno.createUnoStruct("com.sun.star.lang.Locale")
        loc.Language, loc.Country = "zh", "CN"
        rules = self.uno.createUnoStruct("com.sun.star.i18n.ForbiddenCharacters")
        rules.beginLine, rules.endLine = LINE_HEAD_FORBIDDEN, LINE_TAIL_FORBIDDEN
        doc.ForbiddenCharacters.setForbiddenCharacters(loc, rules)
        # 西文/数字后紧跟“%”“～”时（如“NEU%”“20.0～40”），LibreOffice 会把它们
        # 断到下一行行首。仅在排版时插入零宽连接符 U+2060（Excel 文件中的文字不变）。
        # 反过来，“mmol/L↑(<5.18)”这类整串西文无断点，放不下时整串移到下一行，
        # 上一行留下大段空白。在“↑(”“L(”之间加零宽空格 U+200B，允许在“(”前换行。
        glue = re.compile(r"(?<=[A-Za-z0-9#.)])(?=[%～])")
        split = re.compile(r"(?<=[A-Za-z↑↓])(?=\()")
        for col in (7, 9):
            for r in range(FIRST_DATA_ROW - 1, n_rows):
                cell = sheet.getCellByPosition(col, r)
                text = split.sub("\u200b", glue.sub("\u2060", cell.String))
                if text != cell.String:
                    cell.String = text
        return doc, sheet

    def measure_heights(self, path, first, last):
        """返回 first..last 行在 LibreOffice 中的最适行高（pt）。"""
        doc, sheet = self.open(path, last)
        rows = sheet.getRows()
        heights = []
        for r in range(first, last + 1):
            row = rows.getByIndex(r - 1)
            row.OptimalHeight = True
            heights.append(row.Height / 100 / 25.4 * 72)     # 1/100 mm → pt
        doc.close(True)
        return heights

    def export_pdf(self, path, pdf_path, n_rows):
        doc, _ = self.open(path, n_rows)
        url = self.uno.systemPathToFileUrl(str(Path(pdf_path).resolve()))
        doc.storeToURL(url, (self._prop("FilterName", "calc_pdf_Export"),))
        doc.close(True)


def plan_pages(heights, cap_first, cap_other):
    """按原顺序分页，每页 1–2 人，页数最少；同一人不跨页。"""
    n = len(heights)
    INF = 10 ** 9
    best = [INF] * (n + 1)
    choice = [0] * (n + 1)
    best[0] = 0
    for i in range(1, n + 1):
        for k in (1, 2):
            if i - k < 0:
                continue
            cap = cap_first if i - k == 0 else cap_other
            if sum(heights[i - k:i]) <= cap and best[i - k] + 1 < best[i]:
                best[i] = best[i - k] + 1
                choice[i] = k
    assert best[n] < INF, "存在单人内容超过一页的情况"
    pages = []
    i = n
    while i > 0:
        k = choice[i]
        pages.append(list(range(i - k, i)))
        i -= k
    return pages[::-1]


def main(src, outdir, stem=None):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    annex, title, info, headers, people = read_source(src)
    stem = stem or f"{annex}_{title}"

    cap_other = PAGE_H_PT - HEADER_H - SAFETY_PT
    cap_first = cap_other - sum(TITLE_ROWS.values())
    limit = min(MAX_ROW_PT, cap_other)

    last = FIRST_DATA_ROW + len(people) - 1
    with tempfile.TemporaryDirectory() as tmp, Office() as office:
        fonts = [COLS[7][3]] * len(people)
        while True:
            probe = Path(tmp) / "measure.xlsx"
            build_workbook(annex, title, info, people, fonts=fonts).save(probe)
            measured = office.measure_heights(probe, FIRST_DATA_ROW, last)
            heights = [round_px(max(h + ROW_PAD_PT, excel_height_estimate(p["vals"], f)))
                       for h, p, f in zip(measured, people, fonts)]
            too_big = [i for i, h in enumerate(heights) if h > limit]
            if not too_big:
                break
            for i in too_big:
                assert fonts[i] - 0.5 >= MIN_LONG_FONT, f"第 {i + 1} 人内容过多"
                fonts[i] -= 0.5

        pages = plan_pages(heights, cap_first, cap_other)
        breaks = [FIRST_DATA_ROW + pg[-1] for pg in pages[:-1]]
        wb = build_workbook(annex, title, info, people, heights=heights, fonts=fonts, breaks=breaks)
        xlsx = outdir / f"{stem}.xlsx"
        wb.save(xlsx)
        pdf = outdir / f"{stem}.pdf"
        office.export_pdf(xlsx, pdf, last)

    plan = {
        "people": len(people),
        "pages": [[people[i]["vals"][0] for i in pg] for pg in pages],
        "heights": heights,
        "lo_measured": [round(h, 2) for h in measured],
        "excel_estimate": [round(excel_height_estimate(p["vals"], f), 2) for p, f in zip(people, fonts)],
        "fonts": fonts,
        "src_rows": [p["src_row"] for p in people],
    }
    (outdir / "layout_plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"人员 {len(people)} 人，共 {len(pages)} 页；"
          f"单人页 {sum(len(p) == 1 for p in pages)}，双人页 {sum(len(p) == 2 for p in pages)}；"
          f"长文本字号 {sorted(set(fonts))}")
    print(xlsx)
    print(pdf)


if __name__ == "__main__":
    main(*sys.argv[1:4])
