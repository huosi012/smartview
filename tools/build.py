#!/usr/bin/env python3
"""根据 config.js 生成两个成品，放在 output/ 目录：

  1. <标题>-交互版.pptx   PowerPoint 交互页（平滑切换：照片从海报原位飞到屏幕中间放大）
  2. <标题>-网页版.html   单文件网页版（图片已内嵌，拷贝这一个文件即可播放）

用法：  python3 tools/build.py
依赖：  pip install python-pptx Pillow
"""
import copy
import base64
import io
import json
from pathlib import Path

from lxml import etree
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"

# 舞台 2800×1000 单位 = 14m×5m；PPT 页面 140cm×50cm（1:10），1 单位 = 0.05cm
STAGE_W, STAGE_H = 2800, 1000
EMU_PER_U = 18000
PT_PER_U = EMU_PER_U / 12700

# 放大后的版面，与 index.html 中的 LAYOUT 保持一致
LAYOUT = {"photoMaxH": 840, "photoMaxW": 1000, "gap": 90, "margin": 120, "panelTop": 120, "panelH": 760}

POSTER_MAX_SIDE = 7000   # 底图最长边上限（像素），超过则缩小以控制文件体积
PHOTO_MAX_SIDE = 2000    # 个人照片最长边上限
MORPH_MS = 1000

LATIN_FONT, EA_FONT = "Microsoft YaHei", "微软雅黑"
RED, GOLD, INK, WHITE = "C8161D", "FFD27A", "2A1A1A", "FFFFFF"

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"


# ---------------------------------------------------------------- 配置与图片

def load_config():
    text = (ROOT / "config.js").read_text("utf-8")
    text = "\n".join(l for l in text.splitlines() if not l.strip().startswith("//"))
    return json.loads(text[text.index("=") + 1:text.rindex(";")])


def prepare_image(path, max_side):
    """读取图片：按 EXIF 方向摆正、限制尺寸。返回 (bytes, mime, (宽, 高))。"""
    raw = path.read_bytes()
    im = Image.open(io.BytesIO(raw))
    rotated = im.getexif().get(0x0112, 1) != 1
    if im.format in ("JPEG", "PNG") and not rotated and max(im.size) <= max_side:
        return raw, Image.MIME[im.format], im.size
    im = ImageOps.exif_transpose(im).convert("RGB")
    im.thumbnail((max_side, max_side), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=90)
    return buf.getvalue(), "image/jpeg", im.size


def pct_rect(a):
    return (a["x"] / 100 * STAGE_W, a["y"] / 100 * STAGE_H, a["w"] / 100 * STAGE_W, a["h"] / 100 * STAGE_H)


def layout(aspect):
    h = LAYOUT["photoMaxH"]
    w = h * aspect
    if w > LAYOUT["photoMaxW"]:
        w = LAYOUT["photoMaxW"]
        h = w / aspect
    px, py = (STAGE_W - w) / 2, (STAGE_H - h) / 2
    right_x = px + w + LAYOUT["gap"]
    return {
        "photo": (px, py, w, h),
        "left": (LAYOUT["margin"], px - LAYOUT["gap"] - LAYOUT["margin"]),
        "right": (right_x, LAYOUT["panelTop"], STAGE_W - LAYOUT["margin"] - right_x, LAYOUT["panelH"]),
    }


# ---------------------------------------------------------------- 文字排版估算

def char_em(ch):
    if ch == " ":
        return 0.3
    return 0.6 if ord(ch) < 0x2000 else 1.0


def count_lines(paras, size, box_w):
    cap = box_w / size * 0.97   # 留一点余量给标点避头尾
    total = 0
    for para in paras:
        used, lines = 2.0, 1    # 首行缩进 2 字
        for ch in para:
            w = char_em(ch)
            if used + w > cap:
                lines += 1
                used = w
            else:
                used += w
        total += lines
    return total


def fit_bio_size(paras, box_w, box_h, max_size=60, min_size=28, line=1.6):
    size = max_size
    while size > min_size and count_lines(paras, size, box_w) * size * line > box_h - size * 0.3:
        size -= 2
    return size


def fit_line_size(text, box_w, max_size, min_size, spacing_em):
    width_em = sum(char_em(c) + spacing_em for c in text) or 1
    return max(min_size, min(max_size, box_w / width_em))


# ---------------------------------------------------------------- PPT 底层工具

def E(u):
    return Emu(int(round(u * EMU_PER_U)))


def a_el(tag, **attrs):
    el = etree.Element(qn("a:" + tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def color_el(hex_color, alpha=None):
    c = a_el("srgbClr", val=hex_color)
    if alpha is not None:
        c.append(a_el("alpha", val=int(alpha * 100000)))
    return c


def strip_style(shape):
    """去掉主题默认样式（阴影、边框等），只保留显式设置。"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def set_fill(shape, hex_color, alpha=1.0):
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        old = spPr.find(qn(tag))
        if old is not None:
            spPr.remove(old)
    fill = a_el("solidFill")
    fill.append(color_el(hex_color, alpha if alpha < 1 else None))
    spPr.insert(2, fill)   # xfrm, prstGeom 之后


def set_gradient(shape, hex_color):
    """由透明渐变到实色（从左到右）。"""
    spPr = shape._element.spPr
    old = spPr.find(qn("a:solidFill"))
    grad = a_el("gradFill", rotWithShape=1)
    stops = a_el("gsLst")
    for pos, alpha in ((0, 0.0), (100000, None)):
        gs = a_el("gs", pos=pos)
        gs.append(color_el(hex_color, alpha))
        stops.append(gs)
    grad.append(stops)
    grad.append(a_el("lin", ang=0, scaled=0))
    old.addprevious(grad)
    spPr.remove(old)


def set_no_line(shape):
    shape.line.fill.background()


def set_name(shape, name):
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("name", name)


def add_rect(slide, x, y, w, h, name, color, alpha=1.0, shape_type=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape_type, E(x), E(y), E(w), E(h))
    strip_style(shp)
    set_fill(shp, color, alpha)
    set_no_line(shp)
    set_name(shp, name)
    return shp


def add_hotspot(slide, x, y, w, h, name, target):
    # 100% 透明但有填充的矩形：放映时可点击，看不见
    shp = add_rect(slide, x, y, w, h, name, "000000", alpha=0.0)
    shp.click_action.target_slide = target
    return shp


def style_run(run, size_u, color, bold=False, spacing_em=0.0):
    f = run.font
    size_pt = size_u * PT_PER_U
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = LATIN_FONT
    rPr = run._r.get_or_add_rPr()
    rPr.find(qn("a:latin")).addnext(a_el("ea", typeface=EA_FONT))
    if spacing_em:
        rPr.set("spc", str(int(size_pt * spacing_em * 100)))


def add_text(slide, x, y, w, h, name, paras, size_u, color, bold=False, align=PP_ALIGN.LEFT,
             spacing_em=0.0, line=None, indent_em=0.0, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    set_name(tb, name)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = Pt(size_u * PT_PER_U * line)
        if indent_em:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "0")
            pPr.set("indent", str(int(E(size_u * indent_em))))
        run = p.add_run()
        run.text = text
        style_run(run, size_u, color, bold, spacing_em)
    return tb


def add_crop_picture(slide, image_bytes, rect_pct, x, y, w, h, name):
    """从海报上裁出一块（与海报像素完全重合，看不出接缝）。"""
    pic = slide.shapes.add_picture(io.BytesIO(image_bytes), E(x), E(y), E(w), E(h))
    ax, ay, aw, ah = (rect_pct[k] / 100 for k in ("x", "y", "w", "h"))
    pic.crop_left, pic.crop_top = ax, ay
    pic.crop_right, pic.crop_bottom = 1 - ax - aw, 1 - ay - ah
    round_corners(pic)
    set_name(pic, name)
    return pic


def round_corners(pic, adj=3000):
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    av = pic._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    av.append(a_el("gd", name="adj", fmla="val %d" % adj))


def add_frame_and_shadow(pic, width_u):
    pic.line.color.rgb = RGBColor.from_string(WHITE)
    pic.line.width = E(width_u)
    effects = a_el("effectLst")
    shadow = a_el("outerShdw", blurRad=int(E(120)), dist=int(E(30)), dir=5400000, algn="t", rotWithShape=0)
    shadow.append(color_el("000000", 0.65))
    effects.append(shadow)
    pic._element.spPr.append(effects)


def set_transition(slide):
    """平滑（Morph）切换；旧版 Office / 不支持时退回为淡入淡出。禁止鼠标单击翻页。"""
    xml = (
        f'<mc:AlternateContent xmlns:mc="{NS_MC}">'
        f'<mc:Choice xmlns:p159="{NS_P159}" Requires="p159">'
        f'<p:transition xmlns:p="{NS_P}" xmlns:p14="{NS_P14}" spd="slow" p14:dur="{MORPH_MS}" advClick="0">'
        f'<p159:morph option="byObject"/></p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition xmlns:p="{NS_P}" spd="slow" advClick="0"><p:fade/></p:transition>'
        f'</mc:Fallback></mc:AlternateContent>'
    )
    slide._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(xml))


def enable_loop(prs):
    """放映时循环（最后一位之后回到总览，不会出现黑屏“放映结束”）。"""
    for part in prs.part.package.iter_parts():
        if str(part.partname) == "/ppt/presProps.xml":
            root = etree.fromstring(part.blob)
            show = etree.Element(f"{{{NS_P}}}showPr", loop="1", showNarration="1")
            etree.SubElement(show, f"{{{NS_P}}}present")
            etree.SubElement(show, f"{{{NS_P}}}sldAll")
            ext = root.find(f"{{{NS_P}}}extLst")
            if ext is not None:
                ext.addprevious(show)
            else:
                root.append(show)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---------------------------------------------------------------- 生成 PPT

def build_pptx(cfg, out_path):
    people = cfg["people"]
    poster_bytes, _, _ = prepare_image(ROOT / cfg["background"], POSTER_MAX_SIDE)

    photos = []
    for p in people:
        if p.get("photo"):
            data, _, (w, h) = prepare_image(ROOT / p["photo"], PHOTO_MAX_SIDE)
            photos.append((data, w / h))
        else:
            _, _, cw, ch = pct_rect(p["photoArea"])
            photos.append((None, cw / ch))

    prs = Presentation()
    prs.slide_width, prs.slide_height = E(STAGE_W), E(STAGE_H)
    blank = prs.slide_layouts[6]
    overview = prs.slides.add_slide(blank)
    details = [prs.slides.add_slide(blank) for _ in people]

    def base_layer(slide, enlarged=None):
        slide.shapes.add_picture(io.BytesIO(poster_bytes), 0, 0, E(STAGE_W), E(STAGE_H))
        set_name(slide.shapes[-1], "!!bg")
        for k, p in enumerate(people):
            if k != enlarged:
                add_crop_picture(slide, poster_bytes, p["photoArea"], *pct_rect(p["photoArea"]), name=f"!!p{k + 1}")

    # 第 1 页：总览，每张卡片是一个跳转热区
    base_layer(overview)
    for k, p in enumerate(people):
        add_hotspot(overview, *pct_rect(p.get("hotspot") or p["photoArea"]), name=f"热区{k + 1}", target=details[k])
    set_transition(overview)
    overview.notes_slide.notes_text_frame.text = (
        "总览页。点击照片放大；翻页笔“下一页”按顺序放大；放映中输入页码+回车可直接跳转：\n"
        + "\n".join(f"第 {k + 2} 页：{p['name']}（{p.get('unit', '')}）" for k, p in enumerate(people))
    )

    # 第 2… 页：每人一页
    for k, (slide, p) in enumerate(zip(details, people)):
        image, aspect = photos[k]
        L = layout(aspect)
        base_layer(slide, enlarged=k)
        add_rect(slide, 0, 0, STAGE_W, STAGE_H, "!!dim", "1E0003", alpha=0.82)

        px, py, pw, ph = L["photo"]
        if image is None:
            pic = add_crop_picture(slide, poster_bytes, p["photoArea"], px, py, pw, ph, name=f"!!p{k + 1}")
        else:
            pic = slide.shapes.add_picture(io.BytesIO(image), E(px), E(py), E(pw), E(ph))
            round_corners(pic)
            set_name(pic, f"!!p{k + 1}")
        add_frame_and_shadow(pic, 12)

        # 左侧：标语 / 姓名 / 金线 / 单位（右对齐，贴着照片）
        lx, lw = L["left"]
        right_edge = lx + lw
        name_size = fit_line_size(p["name"], lw, 180, 70, 0.12)
        unit_size = fit_line_size(p.get("unit", ""), lw, 76, 36, 0.1)
        add_text(slide, lx, 273, lw + 46 * 0.3, 60, f"标语{k + 1}", [cfg.get("tagline", "")], 46, GOLD,
                 align=PP_ALIGN.RIGHT, spacing_em=0.3)
        add_text(slide, lx, 333 + (180 - name_size) * 0.6, lw + name_size * 0.12, name_size * 1.25,
                 f"姓名{k + 1}", [p["name"]], name_size, WHITE, bold=True, align=PP_ALIGN.RIGHT, spacing_em=0.12)
        bar = add_rect(slide, right_edge - 300, 584, 300, 10, f"金线{k + 1}", GOLD, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        bar.adjustments[0] = 0.5
        set_gradient(bar, GOLD)
        add_text(slide, lx, 628, lw + unit_size * 0.1, 100, f"单位{k + 1}", [p.get("unit", "")], unit_size, GOLD,
                 bold=True, align=PP_ALIGN.RIGHT, spacing_em=0.1)

        # 右侧：个人简介卡片
        rx, ry, rw, rh = L["right"]
        card = add_rect(slide, rx, ry, rw, rh, f"简介卡{k + 1}", WHITE, alpha=0.96, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        card.adjustments[0] = 24 / min(rw, rh)
        pad_x, pad_y = 60, 54
        add_rect(slide, rx + pad_x, ry + pad_y, 14, 56, f"简介标记{k + 1}", RED, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(slide, rx + pad_x + 36, ry + pad_y - 4, rw - pad_x * 2 - 36, 64, f"简介标题{k + 1}",
                 [cfg.get("bioLabel", "个人简介")], 52, RED, bold=True, spacing_em=0.15, anchor=MSO_ANCHOR.MIDDLE)
        bio_paras = [s.strip() for s in str(p.get("bio", "")).split("\n") if s.strip()]
        bx, by = rx + pad_x, ry + pad_y + 56 + 30
        bw, bh = rw - pad_x * 2, rh - pad_y * 2 - 56 - 30
        size = fit_bio_size(bio_paras, bw, bh)
        add_text(slide, bx, by, bw, bh, f"简介{k + 1}", bio_paras, size, INK,
                 align=PP_ALIGN.JUSTIFY, line=1.6, indent_em=2)

        # 点击任意处返回总览
        add_hotspot(slide, 0, 0, STAGE_W, STAGE_H, f"返回{k + 1}", target=overview)
        set_transition(slide)
        slide.notes_slide.notes_text_frame.text = (
            f"第 {k + 1} 位：{p['name']}（{p.get('unit', '')}）\n点击任意处返回总览；翻页笔“下一页”放大下一位。"
        )

    enable_loop(prs)
    prs.save(out_path)


# ---------------------------------------------------------------- 生成单文件网页

def data_uri(path, max_side):
    data, mime, _ = prepare_image(path, max_side)
    return f"data:{mime};base64," + base64.b64encode(data).decode()


def build_html(cfg, out_path):
    html = (ROOT / "index.html").read_text("utf-8")
    c = copy.deepcopy(cfg)
    c["background"] = data_uri(ROOT / cfg["background"], POSTER_MAX_SIDE)
    for p in c["people"]:
        if p.get("photo"):
            p["photo"] = data_uri(ROOT / p["photo"], PHOTO_MAX_SIDE)
    tag = '<script src="config.js"></script>'
    if tag not in html:
        raise SystemExit("index.html 中找不到 config.js 的引用")
    inline = json.dumps(c, ensure_ascii=False).replace("</", "<\\/")
    out_path.write_text(html.replace(tag, f"<script>window.SMARTVIEW_CONFIG = {inline};</script>"), "utf-8")


def main():
    cfg = load_config()
    OUT.mkdir(exist_ok=True)
    title = cfg.get("title", "smartview")
    pptx_path = OUT / f"{title}-交互版.pptx"
    html_path = OUT / f"{title}-网页版.html"
    build_pptx(cfg, pptx_path)
    build_html(cfg, html_path)
    for p in (pptx_path, html_path):
        print(f"已生成 {p.relative_to(ROOT)}  ({p.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
