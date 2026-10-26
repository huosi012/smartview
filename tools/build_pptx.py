#!/usr/bin/env python3
"""生成交互版 PPT：output/<标题>-交互版.pptx（数据来自 config.js）

  · 每页 8 人排成一行（照片 + 名牌 + 下方文字框，同最初海报的版式），人数多就自动分页（24 人 = 3 页）
  · 点照片：照片“平滑”飞到屏幕中间放大，展开人物展板（姓名、单位、一句话事迹、正面行为介绍）；
    点任意处回到本页
  · 翻页笔“下一页”按顺序走：第 1 页 → 第 1～8 位 → 第 2 页 → 第 9～16 位 → … → 回到第 1 页
  · 背景图放在“幻灯片母版”里，换背景只要换一次
  · 页面 140cm × 50cm（= 14m × 5m 的 1:10）

用法：  python3 tools/build_pptx.py
依赖：  pip install python-pptx Pillow
"""
import io
import math

import numpy as np
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from build import ROOT, OUT, load_config

# 舞台 2800×1000 单位 = 14m×5m；PPT 页面 140cm×50cm，1 单位 = 0.05cm
STAGE_W, STAGE_H = 2800, 1000
EMU_PER_U = 18000
PT_PER_U = EMU_PER_U / 12700
PER_PAGE = 8

# 总览页卡片（一行 8 张）：金边 5 + 衬纸 5 包住 4:3 照片，下面是名牌，再下面是文字框
CARD_W, PAD, PLATE, GAP_X = 300, 10, 66, 34
PHOTO_W = CARD_W - 2 * PAD
PHOTO_H = round(PHOTO_W * 3 / 4)
FRAME_H = PAD + PHOTO_H + 5 + PLATE + 5
CARD_TOP = 300                                   # 顶部约 1.5 米留给背景上的大标题
BIO_TOP, BIO_BOTTOM = CARD_TOP + FRAME_H + 14, 935
BUTTON_Y, BUTTON_H = 946, 44                     # 翻页按钮在卡片下方的两个角

# 放大页版面（与网页版一致）：照片连相框居中，展板横贯全屏
LAYOUT = {"photoMaxH": 860, "photoMaxW": 1000, "frame": 28, "boardX": 70, "boardMinH": 600,
          "boardMaxH": 760, "overhang": 50, "gap": 70, "inset": 70, "padY": 50}

MORPH_MS = 1000
PHOTO_MAX_SIDE = 2000
LATIN_FONT, EA_FONT = "Microsoft YaHei", "微软雅黑"
GOLD, CREAM, RED, DARKGOLD = "FFD27A", "FFF3E2", "C8161D", "A06A1C"
GOLD_TEXT = [(0, "FFF8DC"), (46000, "FFE39A"), (100000, "E0A43A")]
GOLD_FRAME = [(0, "FFF4C8"), (22000, "D9A03E"), (46000, "FFF1C0"), (70000, "B9822A"), (100000, "FFE7A8")]

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
CJK_FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"


# ---------------------------------------------------------------- 图片

def jpeg_bytes(im, quality=90):
    buf = io.BytesIO()
    im.convert("RGB").save(buf, "JPEG", quality=quality)
    return buf.getvalue()


def load_photo(rel):
    """读照片：按 EXIF 摆正、限制尺寸。返回 (jpeg 字节, 宽/高)。"""
    im = ImageOps.exif_transpose(Image.open(ROOT / rel))
    im.thumbnail((PHOTO_MAX_SIDE, PHOTO_MAX_SIDE), Image.LANCZOS)
    return jpeg_bytes(im), im.width / im.height


def placeholder_photo():
    """没照片的人：米金底 + 人像剪影 + “照片待补充”。"""
    w, h = 1200, 900
    im = Image.new("RGB", (w, h))
    d = ImageDraw.Draw(im)
    for y in range(h):
        t = y / h
        d.line([0, y, w, y], fill=(int(251 - 25 * t), int(236 - 47 * t), int(208 - 82 * t)))
    d.ellipse([w / 2 - 120, 200, w / 2 + 120, 440], fill=(255, 250, 240))
    d.pieslice([w / 2 - 260, 470, w / 2 + 260, 990], 180, 360, fill=(255, 250, 240))
    try:
        font = ImageFont.truetype(CJK_FONT, 64)
        d.text((w / 2, 800), "照片待补充", font=font, fill=(150, 90, 40), anchor="mm")
    except OSError:
        pass
    return jpeg_bytes(im), w / h


def default_background():
    """没有指定背景时：红色渐变 + 光芒 + 底部丝带 + 标志 + “我为核安全文化代言”标题（5600×2000）。"""
    w, h = STAGE_W * 2, STAGE_H * 2
    cx, cy, r = w / 2, -0.1 * h, 1.25 * w
    ys, xs = np.mgrid[0:h, 0:w].astype(np.float32)
    t = np.clip(np.hypot((xs - cx) / r, (ys - cy) / (r * 0.55)), 0, 1)
    stops = [(0, (226, 48, 47)), (0.42, (200, 22, 29)), (1, (138, 9, 16))]
    rgb = np.dstack([np.interp(t, [s[0] for s in stops], [s[1][c] for s in stops]) for c in range(3)])
    base = Image.fromarray(rgb.astype(np.uint8), "RGB")
    over = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(over)
    for k in range(48):                                   # 顶部光芒
        a = math.radians(-90 + (k - 24) * 3.6)
        d.polygon([(cx, -600), (cx + math.cos(a - 0.012) * 4000, -600 + math.sin(a - 0.012) * -4000),
                   (cx + math.cos(a + 0.012) * 4000, -600 + math.sin(a + 0.012) * -4000)], fill=(255, 255, 255, 10))
    for k, (y0, amp, alpha) in enumerate([(1760, 120, 46), (1860, 90, 60)]):   # 底部丝带
        pts = [(x, y0 + amp * math.sin(x / w * math.pi * 2.2 + k)) for x in range(0, w + 40, 40)]
        d.polygon(pts + [(w, h), (0, h)], fill=(255, 225, 170, alpha))
    over = over.filter(ImageFilter.GaussianBlur(6))
    base = Image.alpha_composite(base.convert("RGBA"), over)
    for name, x, y, height in (("assets/logo.png", 140, 92, 148), ("assets/title.png", None, 104, 172)):
        path = ROOT / name
        if path.exists():
            im = Image.open(path).convert("RGBA")
            im = im.resize((round(im.width * height / im.height), height), Image.LANCZOS)
            base.alpha_composite(im, ((w - im.width) // 2 if x is None else x, y))
    return jpeg_bytes(base, 88)


# ---------------------------------------------------------------- 文字排版估算

def char_em(ch):
    if ch == " ":
        return 0.3
    return 0.58 if ord(ch) < 0x2000 else 1.0


def text_em(text, spacing=0.0):
    return sum(char_em(c) + spacing for c in text)


def count_lines(paras, size, box_w, indent=0.0):
    cap = box_w / size * 0.97
    total = 0
    for para in paras:
        used, lines = indent, 1
        for ch in para:
            w = char_em(ch)
            if used + w > cap:
                lines, used = lines + 1, w
            else:
                used += w
        total += lines
    return total


def fit_size(paras, box_w, box_h, max_size, min_size, line, indent=0.0):
    size = max_size
    while size > min_size and count_lines(paras, size, box_w, indent) * size * line > box_h - size * 0.3:
        size -= 2
    return size


# ---------------------------------------------------------------- PPT 底层工具

def E(u):
    return Emu(int(round(u * EMU_PER_U)))


def a_el(tag, **attrs):
    el = etree.Element(qn("a:" + tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def color_el(hex_color, alpha=None):
    c = a_el("srgbClr", val=hex_color)
    if alpha is not None:
        c.append(a_el("alpha", val=int(alpha * 100000)))
    return c


def grad_el(stops, angle=90, alpha=None):
    g = a_el("gradFill", rotWithShape=1)
    lst = a_el("gsLst")
    for pos, col in stops:
        gs = a_el("gs", pos=pos)
        gs.append(color_el(col, alpha))
        lst.append(gs)
    g.append(lst)
    g.append(a_el("lin", ang=int(angle * 60000), scaled=0))
    return g


def set_name(shape, name):
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("name", name)


def replace_fill(spPr, fill):
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        old = spPr.find(qn(tag))
        if old is not None:
            spPr.remove(old)
    spPr.find(qn("a:prstGeom")).addnext(fill)


def add_shape(slide, x, y, w, h, name, fill=None, alpha=1.0, line=None, line_w=0, radius=0, shape=None, angle=90):
    """矩形 / 圆角矩形。fill：颜色字符串或渐变 [(pos, 颜色)]（angle 为渐变方向，90 = 从上到下）；line：边线颜色。"""
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)                       # 去掉主题默认的阴影、边框
    spPr = shp._element.spPr
    if fill is None:
        replace_fill(spPr, a_el("noFill"))
    elif isinstance(fill, str):
        f = a_el("solidFill")
        f.append(color_el(fill, alpha if alpha < 1 else None))
        replace_fill(spPr, f)
    else:
        replace_fill(spPr, grad_el(fill, angle, alpha=alpha if alpha < 1 else None))
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = E(line_w)
    else:
        shp.line.fill.background()
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    set_name(shp, name)
    return shp


def add_hotspot(slide, x, y, w, h, name, target):
    shp = add_shape(slide, x, y, w, h, name, fill="000000", alpha=0.0)   # 全透明但有填充：放映时可点
    shp.click_action.target_slide = target
    return shp


def add_diamond(slide, cx, cy, size, name, color=GOLD):
    return add_shape(slide, cx - size / 2, cy - size / 2, size, size, name, fill=color, shape=MSO_SHAPE.DIAMOND)


def style_run(run, size_u, color, bold=False, spacing_em=0.0, gradient=None, shadow=False):
    f = run.font
    size_pt = size_u * PT_PER_U
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = LATIN_FONT
    rPr = run._r.get_or_add_rPr()
    rPr.find(qn("a:latin")).addnext(a_el("ea", typeface=EA_FONT))
    if spacing_em:
        rPr.set("spc", str(int(size_pt * spacing_em * 100)))
    if gradient:                                      # 金色渐变字
        solid = rPr.find(qn("a:solidFill"))
        solid.addprevious(grad_el(gradient, 90))
        rPr.remove(solid)
    if shadow:
        eff = a_el("effectLst")
        sh = a_el("outerShdw", blurRad=int(E(16)), dist=int(E(6)), dir=5400000, algn="t", rotWithShape=0)
        sh.append(color_el("000000", 0.5))
        eff.append(sh)
        rPr.find(qn("a:latin")).addprevious(eff)


def add_text(slide, x, y, w, h, name, paras, size_u, color, bold=False, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, spacing_em=0.0, line=None, indent_em=0.0, gradient=None, shadow=False):
    """paras：字符串列表；也可以是 [(文字, 字号, 颜色, 粗体[, 首行缩进])] 的列表，每段不同样式。"""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    set_name(tb, name)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        text, sz, col, bd = item[:4] if isinstance(item, tuple) else (item, size_u, color, bold)
        ind = item[4] if isinstance(item, tuple) and len(item) > 4 else indent_em
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = Pt(sz * PT_PER_U * line)
        if ind:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "0")
            pPr.set("indent", str(int(E(sz * ind))))
        run = p.add_run()
        run.text = text
        style_run(run, sz, col, bd, spacing_em, gradient, shadow)
    return tb


def add_picture(slide, data, x, y, w, h, name, crop=None, line=None, line_w=0):
    pic = slide.shapes.add_picture(io.BytesIO(data), E(x), E(y), E(w), E(h))
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if line:
        pic.line.color.rgb = RGBColor.from_string(line)
        pic.line.width = E(line_w)
    set_name(pic, name)
    return pic


def cover_crop(aspect, target):
    """把照片裁成 target 比例（铺满不变形）；偏高的照片多裁下面，保住脸。"""
    if aspect > target:
        side = (1 - target / aspect) / 2
        return (side, 0, side, 0)
    cut = 1 - aspect / target
    return (0, cut * 0.3, 0, cut * 0.7)


def set_transition(slide):
    """平滑（Morph）切换；旧版 Office 退回淡入淡出。禁止鼠标单击翻页（只能点热区或用翻页笔）。"""
    xml = (
        f'<mc:AlternateContent xmlns:mc="{NS_MC}">'
        f'<mc:Choice xmlns:p159="{NS_P159}" Requires="p159">'
        f'<p:transition xmlns:p="{NS_P}" xmlns:p14="{NS_P14}" spd="slow" p14:dur="{MORPH_MS}" advClick="0">'
        f'<p159:morph option="byObject"/></p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition xmlns:p="{NS_P}" spd="slow" advClick="0"><p:fade/></p:transition>'
        f'</mc:Fallback></mc:AlternateContent>'
    )
    slide._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(xml))


def set_master_background(prs, image_bytes):
    """背景图放进幻灯片母版：所有页面共用，换背景只改一处。"""
    master = prs.slide_master
    _, rId = master.part.get_or_add_image_part(io.BytesIO(image_bytes))
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rId}"/><a:srcRect/>'
        f'<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    cSld = master._element.find(qn("p:cSld"))
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    cSld.insert(0, bg)


def enable_loop(prs):
    """循环放映：放完回到第 1 页，不会出现黑屏“放映结束”。"""
    for part in prs.part.package.iter_parts():
        if str(part.partname) == "/ppt/presProps.xml":
            root = etree.fromstring(part.blob)
            show = etree.Element(f"{{{NS_P}}}showPr", loop="1", showNarration="1")
            etree.SubElement(show, f"{{{NS_P}}}present")
            etree.SubElement(show, f"{{{NS_P}}}sldAll")
            ext = root.find(f"{{{NS_P}}}extLst")
            if ext is not None:
                ext.addprevious(show)
            else:
                root.append(show)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---------------------------------------------------------------- 总览页

def card_x(k):
    """本页第 k 张卡片（0～7）的左边位置：一行居中排开。"""
    return (STAGE_W - (PER_PAGE * CARD_W + (PER_PAGE - 1) * GAP_X)) / 2 + k * (CARD_W + GAP_X)


def build_overview(slide, page, pages, members, details, overviews):
    for k, (idx, person, photo) in enumerate(members):
        x, y = card_x(k), CARD_TOP
        n = idx + 1
        # 照片 + 名牌
        add_shape(slide, x, y, CARD_W, FRAME_H, f"卡框{n}", fill=GOLD_FRAME, radius=12, angle=45)
        add_shape(slide, x + 5, y + 5, CARD_W - 10, PHOTO_H + 10, f"衬纸{n}", fill="FFFAF0")
        data, aspect = photo
        add_picture(slide, data, x + PAD, y + PAD, PHOTO_W, PHOTO_H, f"!!p{n}", crop=cover_crop(aspect, 4 / 3))
        py = y + PAD + PHOTO_H + 5
        add_shape(slide, x + 5, py, CARD_W - 10, PLATE, f"名牌底{n}", fill=[(0, "C81A21"), (100000, "8C0910")])
        add_shape(slide, x + 5, py, CARD_W - 10, 3, f"名牌金线{n}", fill="FFDC8C")
        name_size = min(32, (CARD_W - 30) / text_em(person["name"], 0.08))
        unit_size = min(19, (CARD_W - 30) / max(1, text_em(person.get("unit", ""))))
        add_text(slide, x + 5, py + 3, CARD_W - 10, PLATE - 3, f"名牌{n}",
                 [(person["name"], name_size, GOLD, True), (person.get("unit", ""), unit_size, "FFF1D6", False)],
                 32, GOLD)

        # 下方文字框：一句话事迹（红色粗体）+ 正面行为介绍
        bh = BIO_BOTTOM - BIO_TOP
        add_shape(slide, x, BIO_TOP, CARD_W, bh, f"文字框{n}", fill=[(0, "FFFDF7"), (100000, "FBF1DC")],
                  line="E9C46A", line_w=2, radius=12)
        add_shape(slide, x + 12, BIO_BOTTOM - 14, CARD_W - 24, 5, f"文字框饰条{n}",
                  fill=[(0, "C8161D"), (60000, "E0A43A"), (100000, "FFD27A")], angle=0)
        hl = person.get("highlight", "").strip()
        bio = [t.strip() for t in str(person.get("bio", "")).split("\n") if t.strip()]
        tw, th = CARD_W - 32, bh - 14 - 26
        short = not hl and sum(len(t) for t in bio) <= 30
        size = 22
        while size > 12 and ((count_lines([hl], size, tw) if hl else 0) +
                             count_lines(bio, size, tw, 0 if short else 2)) * size * 1.5 > th - size * 0.3:
            size -= 1
        paras = ([("“" + hl + "”", size, "B8121A", True, 0)] if hl else []) + \
                [(t, size, "3A2A1A", False, 0 if short else 2) for t in bio]
        add_text(slide, x + 16, BIO_TOP + 14, tw, th, f"文字{n}", paras, size, "3A2A1A",
                 align=PP_ALIGN.CENTER if short else PP_ALIGN.JUSTIFY,
                 anchor=MSO_ANCHOR.MIDDLE if short else MSO_ANCHOR.TOP, line=1.5)
        add_hotspot(slide, x, y, CARD_W, BIO_BOTTOM - y, f"热区{n}", details[idx])

    # 翻页按钮（卡片下方左右两角）
    def button(text, x, target, name):
        add_shape(slide, x, BUTTON_Y, 170, BUTTON_H, name + "底", fill="3A0306", alpha=0.55, line=GOLD, line_w=2,
                  radius=BUTTON_H / 2)
        add_text(slide, x, BUTTON_Y, 170, BUTTON_H, name + "字", [text], 24, GOLD, spacing_em=0.1)
        add_hotspot(slide, x, BUTTON_Y, 170, BUTTON_H, name, target)
    if page > 0:
        button("‹ 上一页", 40, overviews[page - 1], "上一页")
    if page < pages - 1:
        button("下一页 ›", STAGE_W - 210, overviews[page + 1], "下一页")
    set_transition(slide)


# ---------------------------------------------------------------- 放大页（人物展板）

def layout(aspect):
    F = LAYOUT["frame"]
    ih = LAYOUT["photoMaxH"] - 2 * F
    iw = ih * aspect
    if iw > LAYOUT["photoMaxW"] - 2 * F:
        iw = LAYOUT["photoMaxW"] - 2 * F
        ih = iw / aspect
    w, h = iw + 2 * F, ih + 2 * F
    photo = ((STAGE_W - w) / 2, (STAGE_H - h) / 2, w, h)
    bh = min(max(h - 2 * LAYOUT["overhang"], LAYOUT["boardMinH"]), LAYOUT["boardMaxH"])
    board = (LAYOUT["boardX"], (STAGE_H - bh) / 2, STAGE_W - 2 * LAYOUT["boardX"], bh)
    top, height = board[1] + LAYOUT["padY"], bh - 2 * LAYOUT["padY"]
    left_x, right_x = board[0] + LAYOUT["inset"], photo[0] + w + LAYOUT["gap"]
    return {"photo": photo, "board": board,
            "left": (left_x, top, photo[0] - LAYOUT["gap"] - left_x, height),
            "right": (right_x, top, board[0] + board[2] - LAYOUT["inset"] - right_x, height)}


def gold_rule(slide, x, y, w, name, flip=False):
    """渐隐的金线（一端透明）。"""
    stops = [(0, GOLD), (100000, GOLD)]
    shp = add_shape(slide, x, y, w, 2.5, name, fill=stops, angle=0)
    grad = shp._element.spPr.find(qn("a:gradFill"))
    first = grad.find(qn("a:gsLst"))[1 if flip else 0].find(qn("a:srgbClr"))
    first.append(a_el("alpha", val=0))
    return shp


def build_detail(slide, idx, person, photo, cfg, back):
    n = idx + 1
    data, aspect = photo
    L = layout(aspect)

    add_shape(slide, 0, 0, STAGE_W, STAGE_H, "!!dim", fill="1E0003", alpha=0.82)

    # 展板：深红底 + 双金线 + 四角纹饰
    bx, by, bw, bh = L["board"]
    add_shape(slide, bx, by, bw, bh, "!!board", fill=[(0, "A3121A"), (58000, "7C0A11"), (100000, "5A050A")],
              line="FFD682", line_w=2.5, radius=30)
    add_shape(slide, bx + 18, by + 18, bw - 36, bh - 36, "!!boardLine", fill=None, line="C9A260", line_w=1.5, radius=20)
    for cx, cy, sx, sy, tag in ((bx, by, 1, 1, "tl"), (bx + bw, by, -1, 1, "tr"), (bx, by + bh, 1, -1, "bl"), (bx + bw, by + bh, -1, -1, "br")):
        ox, oy = cx + sx * 8, cy + sy * 8
        add_shape(slide, ox if sx > 0 else ox - 80, oy if sy > 0 else oy - 4, 80, 4, f"!!角横{tag}", fill=GOLD)
        add_shape(slide, ox if sx > 0 else ox - 4, oy if sy > 0 else oy - 80, 4, 80, f"!!角竖{tag}", fill=GOLD)
        add_diamond(slide, cx + sx * 34, cy + sy * 34, 16, f"!!角菱{tag}")

    # 照片：金框 + 米白衬纸 + 照片（细金线）
    px, py, pw, ph = L["photo"]
    F = LAYOUT["frame"]
    add_shape(slide, px - 3, py - 3, pw + 6, ph + 6, "!!frameEdge", fill="7A4A0C", radius=18)
    add_shape(slide, px, py, pw, ph, "!!frame", fill=GOLD_FRAME, radius=16, angle=45)
    add_shape(slide, px + 14, py + 14, pw - 28, ph - 28, "!!mat", fill=[(0, "FFFDF6"), (100000, "F6EAD0")],
              line="96641E", line_w=2)
    add_picture(slide, data, px + F, py + F, pw - 2 * F, ph - 2 * F, f"!!p{n}", line="C99A3C", line_w=2)
    for cx, cy, tag in ((px, py, "tl"), (px + pw, py, "tr"), (px, py + ph, "bl"), (px + pw, py + ph, "br")):
        add_diamond(slide, cx, cy, 26, f"!!框菱{tag}", color=RED)

    # 左侧：标语 / 姓名 / 分隔 / 单位 / 一句话事迹（居中排列，整体上下居中）
    lx, ly, lw, lh = L["left"]
    cxl = lx + lw / 2
    name = person["name"]
    name_size = min(160, lw / (text_em(name, 0.14) + 0.14))
    unit = person.get("unit", "")
    unit_size = min(56, (lw * 0.92 - 116) / max(1, text_em(unit, 0.14)))
    hl = person.get("highlight", "").strip()
    hl_w = lw - 32
    hl_size = fit_size(["“" + hl + "”"], hl_w, 42 * 1.5 * 3, 42, 28, 1.5) if hl else 0
    hl_lines = count_lines(["“" + hl + "”"], hl_size, hl_w) if hl else 0
    tag = cfg.get("tagline", "")
    heights = [56 if tag else 0, 14 + name_size * 1.22, 68, unit_size * 1.3 + 24 if unit else 0,
               (40 + hl_lines * hl_size * 1.5) if hl else 0]
    y = ly + (lh - sum(heights)) / 2
    if tag:
        tw = text_em(tag, 0.32) * 40
        add_text(slide, cxl - tw / 2, y, tw + 40 * 0.32, 56, f"标语{n}", [tag], 40, GOLD, spacing_em=0.32)
        gold_rule(slide, cxl - tw / 2 - 104, y + 28, 80, f"标语线左{n}")
        gold_rule(slide, cxl + tw / 2 + 24, y + 28, 80, f"标语线右{n}", flip=True)
    y += heights[0] + 14
    add_text(slide, lx - 40, y, lw + 80 + name_size * 0.14, name_size * 1.22, f"姓名{n}", [name], name_size, "FFE39A",
             bold=True, spacing_em=0.14, gradient=GOLD_TEXT, shadow=True)
    y += name_size * 1.22 + 22
    dw = lw * 0.72
    gold_rule(slide, cxl - dw / 2, y + 8, dw / 2 - 18, f"分隔左{n}")
    gold_rule(slide, cxl + 18, y + 8, dw / 2 - 18, f"分隔右{n}", flip=True)
    add_diamond(slide, cxl, y + 9, 22, f"分隔菱{n}")
    y += 46
    if unit:
        uw = min(lw * 0.92, text_em(unit, 0.14) * unit_size + 116)
        uh = unit_size * 1.3 + 24
        add_shape(slide, cxl - uw / 2, y, uw, uh, f"单位框{n}", fill="FFD682", alpha=0.16, line="FFD682", line_w=2, radius=uh / 2)
        add_text(slide, cxl - uw / 2, y, uw + unit_size * 0.14, uh, f"单位{n}", [unit], unit_size, "FFF1CF",
                 bold=True, spacing_em=0.14)
        y += uh
    if hl:
        y += 40
        add_text(slide, lx + 16, y, hl_w, hl_lines * hl_size * 1.5 + 10, f"事迹{n}", ["“" + hl + "”"], hl_size,
                 "FFF3D6", bold=True, line=1.5)

    # 右侧：标题 + 正面行为介绍
    rx, ry, rw, rh = L["right"]
    label = cfg.get("bioLabel", "正面行为介绍")
    lw2 = text_em(label, 0.2) * 50
    cxr = rx + rw / 2
    add_text(slide, cxr - lw2 / 2, ry, lw2 + 10, 66, f"标题{n}", [label], 50, GOLD, bold=True, spacing_em=0.2)
    side = (rw - lw2) / 2 - 40
    if side > 20:
        gold_rule(slide, rx, ry + 33, side, f"标题线左{n}")
        gold_rule(slide, cxr + lw2 / 2 + 40, ry + 33, side, f"标题线右{n}", flip=True)
        add_diamond(slide, rx + side + 6, ry + 34, 12, f"标题菱左{n}")
        add_diamond(slide, cxr + lw2 / 2 + 34, ry + 34, 12, f"标题菱右{n}")
    paras = [s.strip() for s in str(person.get("bio", "")).split("\n") if s.strip()]
    top = ry + 66 + 32
    bh2 = ry + rh - top
    short = sum(len(s) for s in paras) <= 30
    size = fit_size(paras, rw, bh2, 56, 28, 1.7, 0 if short else 2)
    add_text(slide, rx, top, rw, bh2, f"介绍{n}", paras, size, CREAM,
             align=PP_ALIGN.CENTER if short else PP_ALIGN.JUSTIFY, line=1.7, indent_em=0 if short else 2)

    add_hotspot(slide, 0, 0, STAGE_W, STAGE_H, f"返回{n}", back)
    set_transition(slide)


# ---------------------------------------------------------------- 主程序

def build_pptx(cfg, out_path):
    people = cfg["people"]
    photos, blank = [], None
    for p in people:
        if p.get("photo") and (ROOT / p["photo"]).exists():
            photos.append(load_photo(p["photo"]))
        else:
            blank = blank or placeholder_photo()
            photos.append(blank)

    prs = Presentation()
    prs.slide_width, prs.slide_height = E(STAGE_W), E(STAGE_H)
    bg = cfg.get("background")
    set_master_background(prs, jpeg_bytes(Image.open(ROOT / bg)) if bg and (ROOT / bg).exists() else default_background())

    blank_layout = prs.slide_layouts[6]
    pages = math.ceil(len(people) / PER_PAGE)
    overviews, details = [], [None] * len(people)
    for page in range(pages):                      # 页面顺序：第 1 页、第 1～8 位、第 2 页、第 9～16 位……
        overviews.append(prs.slides.add_slide(blank_layout))
        for idx in range(page * PER_PAGE, min(len(people), (page + 1) * PER_PAGE)):
            details[idx] = prs.slides.add_slide(blank_layout)

    number = {s.slide_id: i + 1 for i, s in enumerate(prs.slides)}
    for page in range(pages):
        ids = range(page * PER_PAGE, min(len(people), (page + 1) * PER_PAGE))
        members = [(i, people[i], photos[i]) for i in ids]
        build_overview(overviews[page], page, pages, members, details, overviews)
        overviews[page].notes_slide.notes_text_frame.text = (
            f"第 {page + 1} 页（共 {pages} 页）。点照片放大；翻页笔“下一页”按顺序放大；放映中输入页码再按回车可直接跳转：\n" +
            "\n".join(f"第 {number[details[i].slide_id]} 张幻灯片：第 {i + 1} 位 {people[i]['name']}（{people[i].get('unit', '')}）"
                      for i in ids))
        for i in ids:
            build_detail(details[i], i, people[i], photos[i], cfg, overviews[page])
            details[i].notes_slide.notes_text_frame.text = (
                f"第 {i + 1} 位：{people[i]['name']}（{people[i].get('unit', '')}）\n点任意处回到第 {page + 1} 页；翻页笔“下一页”放大下一位。")

    enable_loop(prs)
    prs.save(out_path)


def main():
    cfg = load_config()
    OUT.mkdir(exist_ok=True)
    out_path = OUT / f"{cfg.get('title', 'smartview')}-交互版.pptx"
    build_pptx(cfg, out_path)
    print(f"已生成 {out_path.relative_to(ROOT)}  ({out_path.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
