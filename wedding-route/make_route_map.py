#!/usr/bin/env python3
"""婚车路线图生成器。

路线：中国石油万金塔派出所北加油站（男方）→ 腰周家（女方）
      → 德惠市人民政府 → 中国石油万金塔派出所北加油站（男方）

用法：
  python3 make_route_map.py                       # 只生成 SVG / HTML
  python3 make_route_map.py --fonts <dir> --png   # 同时用 Chromium 渲染 PNG
  python3 make_route_map.py --fonts <dir> --pptx  # 导出文字可编辑的 PPTX（需 python-pptx）

--fonts 指向 `npm i @fontsource/noto-sans-sc @fontsource/zcool-qingke-huangyou
@fontsource/ma-shan-zheng` 后的 node_modules/@fontsource 目录；不指定时使用系统字体。
日期、姓名、时间安排、电话等请直接修改下面的 INFO。
"""
import argparse
import html
import math
import os
import subprocess

# ---------------------------------------------------------------- 可修改的信息
INFO = {
    "date": "2026.10.28",
    "lunar": "农历九月十九 · 星期三",
    "groom": "谢拾雨",
    "bride": "周心语",
    # (时间, 事项, 备注)；备注可留空
    "schedule": [
        ("5:00 – 5:20", "到达万金塔集结扎花", ""),
        ("5:30", "从万金塔出发（男方）", ""),
        ("6:30", "到达腰周家（女方）", ""),
        ("8:00", "从腰周家出发返回万金塔", "途经德惠市人民政府"),
    ],
    "thanks": "感谢各位朋友，请按图示路线行驶",
    "phone": "",  # 留空则不显示联系电话
}

W, H = 1800, 1280

PINK = "#ff2a73"
ORANGE = "#ff8a00"
BLUE = "#1f49ff"
RED_ROAD = "#e8232f"
BLUE_ROAD = "#27a2e6"
ROAD_FILL = "#fffafd"
TEXT = "#4a4a4a"

SANS = "'Noto Sans SC','WenQuanYi Zen Hei',sans-serif"
DISPLAY = "'ZCOOL QingKe HuangYou','Noto Sans SC','WenQuanYi Zen Hei',sans-serif"
BRUSH = "'Ma Shan Zheng','Noto Sans SC','WenQuanYi Zen Hei',serif"

# ---------------------------------------------------------------- 路网
# 城区直接按高德截图复刻：下面 amap() 里的坐标是城区整体截图（1414×1503）上的像素位置，
# 用同一个比例换算到画布；政府附近的拐弯取自更近的一张截图（2000×1116），
# 用 amap_near() 先对齐到整体截图。303 省道西段是一条长直路，按示意画、适当压缩。
CITY_SCALE = 0.62


def amap(x, y):
    """城区整体截图像素 → 画布。截图上 303 省道转向德惠的路口 (180, 42) 对应画布 (1060, 415)。"""
    return (round(1060 + CITY_SCALE * (x - 180), 1), round(415 + CITY_SCALE * (y - 42), 1))


def amap_near(x, y):
    """政府附近截图像素 → 画布（两张截图用德惠西站、王三姐绿色生态园两点对齐）。"""
    return amap(610 + 0.527 * (x - 452), 615 + 0.527 * (y - 312))


def path(*pts):
    return [amap(x, y) for x, y in pts]


HOME = (205, 252)          # 新郎家：中国石油万金塔派出所北加油站
P0 = (225, 342)            # 加油站出口接 303 省道
TIANTAI = (760, 386)       # 天台镇
J1 = amap(180, 42)         # 303 省道转入 Y037 去德惠
J2 = amap(412, 395)        # 德惠西北路口：去程、回程在这里分开
J3 = amap(460, 925)        # 德农路路口
J4 = amap(552, 1211)       # 德环路 × 新华街 × 德农环路
BRIDE = amap(430, 1340)    # 新娘家：腰周家
S1 = amap_near(933, 858)   # 新华街 × 松柏路（红旗大市场）
S2 = amap_near(988, 886)   # 松柏路上右转一小段后左转向北
S3 = amap_near(1071, 700)
S4 = amap_near(1437, 861)
GOV = amap_near(1460, 952)  # 德惠市人民政府：从北面进，出来再往北
NE = amap_near(1583, 723)

S303 = [(60, 322), P0, (470, 355), TIANTAI, J1]
S303_EAST = path((180, 42), (330, 62), (500, 70), (842, 77), (881, 88), (942, 154),
                 (996, 208), (1112, 323), (1158, 377), (1188, 385), (1412, 431))  # 接北环城路
WEST_RING = path((412, 395), (430, 445), (478, 490), (552, 610), (560, 655), (595, 673),
                 (460, 925), (498, 1140), (510, 1170), (530, 1190), (552, 1211))  # 经德惠西站、德环路
BRIDE_ROAD = path((470, 1292)) + [BRIDE]                 # 交叉口往西南去腰周家的乡道
XINHUA = path((390, 1470), (395, 1430), (405, 1390)) + BRIDE_ROAD[::-1] + [J4, S1]
SONGBAI = path((468, 731), (722, 842)) + [S1, S2] + path((1105, 1000), (1185, 1048))
CITY = [S2, S3, S4, GOV, NE, amap_near(1500, 640)] + path(
    (1081, 700), (742, 560), (696, 462), (620, 424), (600, 432), (558, 430), (528, 407),
    (412, 395))                                           # 回程：政府往北，再往西北出城
DENONG = path((0, 1000), (270, 970), (360, 948), (460, 925))
DENONG_RING = path((552, 1211), (600, 1213), (700, 1262), (800, 1290), (900, 1286),
                   (960, 1272), (1000, 1280))             # 德农环路东段，到 G102 为止
G102 = path((980, 1503), (1000, 1280), (1050, 1200), (1100, 1130), (1185, 1048),
            (1300, 990), (1414, 905))
G1 = path((719, 0), (696, 31), (650, 100), (600, 160), (570, 220), (550, 290), (530, 370),
          (520, 420), (510, 480), (500, 540), (487, 600), (468, 677), (433, 754), (360, 892),
          (300, 1000), (250, 1150), (180, 1380), (140, 1503))   # 京哈高速
RAIL_HS = path((873, 0), (750, 269), (650, 477), (568, 646), (525, 800), (398, 1138),
               (300, 1450), (285, 1503))                  # 高铁（经德惠西站）
RAIL = path((1388, 0), (1288, 223), (885, 1050), (700, 1450), (680, 1503))   # 京哈铁路

MAJOR = (RED_ROAD, 34)         # 303 省道、Y037
STREET = (RED_ROAD, 26)        # 城区主路、新华街、G102
MINOR = (BLUE_ROAD, 26)        # 德环路一线
SIDE = (BLUE_ROAD, 22)         # 不走的次要道路
HIGHWAY = ("#7f93c4", 22)
# (道路, 样式, 渐隐遮罩)：先画的在下面
ROADS = [
    (G1, HIGHWAY, "fadeTop"),
    (S303_EAST, STREET, None),
    (G102, STREET, None),
    (DENONG, SIDE, "fadeLeft"),
    (DENONG_RING, SIDE, None),
    (SONGBAI, STREET, None),
    (S303, MAJOR, None),
    ([J1, J2], MAJOR, None),
    (XINHUA, STREET, None),
    (CITY, STREET, None),
    (WEST_RING, MINOR, None),
    ([HOME, P0], MINOR, None),                           # 加油站出口
]

LEG_TO_BRIDE = [HOME, P0, (470, 355), TIANTAI, J1] + WEST_RING + BRIDE_ROAD
LEG_TO_GOV = BRIDE_ROAD[::-1] + [J4, S1, S2, S3, S4, GOV]
LEG_HOME = CITY[CITY.index(GOV):] + [J1, TIANTAI, (470, 355), P0, HOME]
LEGS = [(LEG_TO_BRIDE, PINK, 30), (LEG_TO_GOV, ORANGE, 30), (LEG_HOME, BLUE, 60)]


# ---------------------------------------------------------------- 几何工具
def seg_lengths(pts):
    return [math.dist(a, b) for a, b in zip(pts, pts[1:])]


def point_at(pts, s):
    """返回折线上弧长 s 处的点、单位切向量，以及到最近折点的距离。"""
    acc = 0.0
    lens = seg_lengths(pts)
    for i, ln in enumerate(lens):
        if acc + ln >= s or i == len(lens) - 1:
            t = (s - acc) / ln
            (x0, y0), (x1, y1) = pts[i], pts[i + 1]
            dx, dy = (x1 - x0) / ln, (y1 - y0) / ln
            near = min(s - acc, acc + ln - s)
            return (x0 + t * (x1 - x0), y0 + t * (y1 - y0)), (dx, dy), near
        acc += ln


def pts_attr(pts):
    return " ".join(f"{x:.1f},{y:.1f}" for x, y in pts)


# ---------------------------------------------------------------- SVG 片段
_text_sink = None   # 导出 PPTX 时收集可编辑文字，此时 SVG 里不画这些文字


def text(x, y, s, size, fill=TEXT, family=SANS, weight=400, anchor="middle",
         rotate=0, halo=None, spacing=0, extra="", editable=True):
    if editable and _text_sink is not None:
        _text_sink.append(dict(x=x, y=y, s=s, size=size, fill=fill, family=family,
                               weight=weight, anchor=anchor, rotate=rotate, halo=halo,
                               spacing=spacing))
        return ""
    if spacing:
        extra += f' letter-spacing="{spacing}"'
    tr = f' transform="rotate({rotate:.1f} {x:.1f} {y:.1f})"' if rotate else ""
    st = (f' stroke="{halo}" stroke-width="{max(4, size // 6)}" stroke-linejoin="round"'
          f' paint-order="stroke"') if halo else ""
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-size="{size}" fill="{fill}" '
            f'font-family="{family}" font-weight="{weight}" text-anchor="{anchor}"'
            f'{tr}{st}{extra}>{html.escape(s)}</text>')


def arrow(x, y, dx, dy, color, length=38, shaft=6, head=16, half=7):
    tx, ty = x - dx * length / 2, y - dy * length / 2
    hx, hy = x + dx * length / 2, y + dy * length / 2
    bx, by = hx - dx * head, hy - dy * head
    nx, ny = -dy, dx
    tri = [(hx, hy), (bx + nx * half, by + ny * half), (bx - nx * half, by - ny * half)]
    return (f'<line x1="{tx:.1f}" y1="{ty:.1f}" x2="{bx + dx * 2:.1f}" y2="{by + dy * 2:.1f}" '
            f'stroke="{color}" stroke-width="{shaft}" stroke-linecap="round"/>'
            f'<polygon points="{pts_attr(tri)}" fill="{color}" stroke="{color}" '
            f'stroke-width="2" stroke-linejoin="round"/>')


def leg_arrows(pts, color, phase, spacing=112, offset=7, margin=40, corner=22):
    out = []
    total = sum(seg_lengths(pts))
    s = margin + phase
    while s < total - margin:
        (x, y), (dx, dy), near = point_at(pts, s)
        if near < corner:
            s += 12
            continue
        # 靠右行驶：箭头画在行进方向右侧车道
        out.append(arrow(x - dy * offset, y + dx * offset, dx, dy, color))
        s += spacing
    return out


def along(a, b, t=0.5, gap=0, side=1):
    """线段 a→b 上比例 t 处、向右（side=1）或向左（side=-1）偏 gap 的点，及便于阅读的角度。"""
    (x0, y0), (x1, y1) = a, b
    ln = math.dist(a, b)
    dx, dy = (x1 - x0) / ln, (y1 - y0) / ln
    x, y = x0 + (x1 - x0) * t - dy * gap * side, y0 + (y1 - y0) * t + dx * gap * side
    ang = math.degrees(math.atan2(dy, dx))
    if ang > 90:
        ang -= 180
    elif ang <= -90:
        ang += 180
    return x, y, ang


def road_label(a, b, s, side=1, gap=40, size=24, t=0.5, fill="#5b5b5b"):
    x, y, ang = along(a, b, t, gap, side)
    return text(x, y + size * 0.35, s, size, fill, SANS, 500, rotate=ang,
                halo="#fff3f7")


def shield(x, y, label, bg, fg, rotate=0, border="#ffffff"):
    w = 14 * len(label) + 16
    tr = f' transform="rotate({rotate:.1f} {x} {y})"' if rotate else ""
    return (f'<g{tr}><rect x="{x - w / 2:.1f}" y="{y - 15}" width="{w}" height="30" rx="6" '
            f'fill="{bg}" stroke="{border}" stroke-width="3"/>'
            + text(x, y + 8, label, 20, fg, SANS, 700, editable=False) + "</g>")


def poi(x, y, label, color, glyph, size=22):
    """地图上的小地标：彩色圆点图标 + 文字，(x, y) 为文字基线起点。"""
    return (f'<circle cx="{x - 16}" cy="{y - size * 0.35:.1f}" r="12" fill="{color}"/>'
            + text(x - 16, y - size * 0.35 + 5, glyph, 14, "#fff", SANS, 700, editable=False)
            + text(x, y, label, size, color, SANS, 500, anchor="start", halo="#fff3f7"))


def pin(x, y, color, inner="♥"):
    body = (f'M{x},{y} C{x - 6},{y - 14} {x - 24},{y - 26} {x - 24},{y - 44} '
            f'A24,24 0 1 1 {x + 24},{y - 44} C{x + 24},{y - 26} {x + 6},{y - 14} {x},{y} Z')
    return (f'<ellipse cx="{x}" cy="{y}" rx="13" ry="5" fill="#000" opacity=".18"/>'
            f'<path d="{body}" fill="{color}" stroke="#fff" stroke-width="4"/>'
            f'<circle cx="{x}" cy="{y - 44}" r="15" fill="#fff"/>'
            + text(x, y - 36, inner, 22, color, SANS, 900, editable=False))


def heart(x, y, size, color, rot=0, opacity=1.0):
    k = size / 60
    return (f'<path transform="translate({x} {y}) rotate({rot}) scale({k:.3f})" '
            f'd="M0,26 C-34,2 -46,-14 -38,-30 C-30,-46 -8,-44 0,-26 '
            f'C8,-44 30,-46 38,-30 C46,-14 34,2 0,26 Z" fill="{color}" opacity="{opacity}"/>')


# ---------------------------------------------------------------- 组装
def build_svg(text_sink=None):
    """text_sink 为列表时，可编辑文字收集进去而不画进 SVG（用作 PPTX 背景）。"""
    global _text_sink
    _text_sink = text_sink
    try:
        return _build_svg()
    finally:
        _text_sink = None


def _build_svg():
    o = []
    o.append(f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" '
             f'viewBox="0 0 {W} {H}">')
    o.append('''<defs>
  <linearGradient id="bg" x1="0" y1="0" x2="0.35" y2="1">
    <stop offset="0" stop-color="#fff7fa"/><stop offset=".55" stop-color="#ffeef4"/>
    <stop offset="1" stop-color="#ffdbe7"/></linearGradient>
  <radialGradient id="blob" cx=".5" cy=".5" r=".5">
    <stop offset="0" stop-color="#ffc2d6" stop-opacity=".85"/>
    <stop offset="1" stop-color="#ffc2d6" stop-opacity="0"/></radialGradient>
  <linearGradient id="seal" x1="0" y1="0" x2="0" y2="1">
    <stop offset="0" stop-color="#f2373f"/><stop offset="1" stop-color="#c8101c"/></linearGradient>
  <filter id="soft" x="-20%" y="-20%" width="140%" height="140%">
    <feDropShadow dx="0" dy="3" stdDeviation="3" flood-color="#c2185b" flood-opacity=".25"/></filter>
  <linearGradient id="fadeTopG" gradientUnits="userSpaceOnUse" x1="0" y1="398" x2="0" y2="446">
    <stop offset="0" stop-color="#000"/><stop offset="1" stop-color="#fff"/></linearGradient>
  <mask id="fadeTop" maskUnits="userSpaceOnUse" x="0" y="0" width="1800" height="1280">
    <rect width="1800" height="1280" fill="url(#fadeTopG)"/></mask>
  <linearGradient id="fadeLeftG" gradientUnits="userSpaceOnUse" x1="948" y1="0" x2="1010" y2="0">
    <stop offset="0" stop-color="#000"/><stop offset="1" stop-color="#fff"/></linearGradient>
  <mask id="fadeLeft" maskUnits="userSpaceOnUse" x="0" y="0" width="1800" height="1280">
    <rect width="1800" height="1280" fill="url(#fadeLeftG)"/></mask>
</defs>''')
    o.append(f'<rect width="{W}" height="{H}" fill="url(#bg)"/>')
    o.append('<circle cx="1730" cy="-40" r="330" fill="url(#blob)"/>')
    o.append('<circle cx="-40" cy="1320" r="380" fill="url(#blob)"/>')
    o.append('<circle cx="1820" cy="1260" r="240" fill="url(#blob)"/>')

    # 铁路（最底层）
    for pts, color in ((RAIL_HS, "#e4b4c6"), (RAIL, "#b7b9cc")):
        o.append(f'<g mask="url(#fadeTop)"><polyline points="{pts_attr(pts)}" fill="none" '
                 f'stroke="{color}" stroke-width="7" stroke-linejoin="round"/>'
                 f'<polyline points="{pts_attr(pts)}" fill="none" stroke="#ffffff" '
                 f'stroke-width="3" stroke-dasharray="14 12"/></g>')

    # 道路：先画所有描边，再画所有填充，交叉口自然连通
    for pts, (color, width), fade in ROADS:
        mask = f' mask="url(#{fade})"' if fade else ""
        o.append(f'<polyline points="{pts_attr(pts)}" fill="none" stroke="{color}" '
                 f'stroke-width="{width + 8}" stroke-linejoin="round" stroke-linecap="butt"{mask}/>')
    for pts, (color, width), fade in ROADS:
        mask = f' mask="url(#{fade})"' if fade else ""
        fill = "#eef1fa" if pts is G1 else ROAD_FILL
        o.append(f'<polyline points="{pts_attr(pts)}" fill="none" stroke="{fill}" '
                 f'stroke-width="{width}" stroke-linejoin="round" stroke-linecap="butt"{mask}/>')

    # 路名、路牌
    o.append(road_label((470, 355), TIANTAI, "303省道", side=-1, gap=38, size=26))
    o.append(shield(452, 314, "S303", "#ffd23f", "#333", rotate=3))
    o.append(road_label(S303_EAST[1], S303_EAST[2], "303省道", side=-1, gap=28, size=24))
    o.append(road_label(S303_EAST[9], S303_EAST[10], "北环城路", side=-1, gap=26, size=22,
                        t=0.45))
    x, y, ang = along(J1, J2, 0.5, 40, -1)
    o.append(shield(x, y, "Y037", "#ffffff", "#333", rotate=ang, border="#333"))
    o.append(shield(*G1[3], "G1", "#1f8f4e", "#fff"))
    o.append(road_label(G1[5], G1[6], "京哈高速", side=-1, gap=26, size=22, fill="#5d6f9c"))
    o.append(road_label(G1[14], G1[15], "京哈高速", side=1, gap=26, size=22, t=0.6,
                        fill="#5d6f9c"))
    o.append(road_label(DENONG[0], DENONG[1], "德农路", side=-1, gap=24, size=22, t=0.62))
    o.append(road_label(J3, WEST_RING[7], "德环路", side=-1, gap=26, size=22, t=0.55))
    o.append(road_label(DENONG_RING[4], DENONG_RING[5], "德农环路", side=-1, gap=24, size=22))
    x, y, ang = along(DENONG_RING[2], DENONG_RING[3], 0.45)
    o.append(shield(x, y, "S001", "#ffd23f", "#333", rotate=ang))
    o.append(road_label(J4, S1, "新华街", side=1, gap=26, size=24))
    o.append(road_label(SONGBAI[0], SONGBAI[1], "松柏路", side=1, gap=24, size=22, t=0.66))
    o.append(road_label(G102[3], G102[4], "G102国道", side=1, gap=26, size=22))
    x, y, ang = along(G102[4], G102[5], 0.55)
    o.append(shield(x, y, "G102", "#e8232f", "#fff", rotate=ang))

    # 地标
    o.append(text(TIANTAI[0], 440, "天台镇", 28, TEXT, SANS, 500))
    o.append(text(70, 292, "万金塔乡", 24, "#777", SANS, 400, anchor="start"))
    o.append(poi(1347, 778, "德惠西站", "#2f6fd6", "站"))
    o.append(poi(1395, 884, "红旗大市场", "#a03cc8", "市"))
    o.append(poi(1300, 968, "龙凤采摘园", "#3a9b3f", "园"))
    o.append(poi(1405, 690, "中林石油加油站", "#2f6fd6", "油"))
    o.append(poi(1580, 748, "王三姐绿色生态园", "#3a9b3f", "园"))

    # 行车箭头
    for pts, color, phase in LEGS:
        o.extend(leg_arrows(pts, color, phase))

    # 地点标记
    o.append(pin(*HOME, BLUE))
    o.append(text(250, 236, "新郎家", 56, BLUE, DISPLAY, 400, anchor="start",
                  halo="#ffffff"))
    o.append(text(252, 278, "中国石油万金塔派出所北加油站", 24, "#2b3f9e", SANS, 500,
                  anchor="start", halo="#fff7fa"))
    o.append(f'<rect x="108" y="196" width="70" height="32" rx="16" fill="{PINK}"/>'
             + text(143, 219, "起·终", 19, "#fff", SANS, 700, editable=False))

    o.append(pin(*BRIDE, PINK))
    o.append(text(1250, 1268, "新娘家", 50, "#e8202c", DISPLAY, 400, anchor="start",
                  halo="#ffffff"))
    o.append(text(1388, 1266, "腰周家", 28, "#c2185b", SANS, 700, anchor="start",
                  halo="#fff3f7"))

    o.append(pin(*GOV, ORANGE, "★"))
    o.append(text(1598, 862, "德惠市", 24, "#d35400", SANS, 900, halo="#ffffff"))
    o.append(text(1598, 892, "人民政府", 24, "#d35400", SANS, 900, halo="#ffffff"))

    # 标题与路线摘要
    o.append(heart(612, 86, 38, PINK, -18, .85))
    o.append(heart(1188, 86, 38, PINK, 18, .85))
    o.append(text(900, 136, "婚车路线图", 104, "#e8202c", DISPLAY, 400, halo="#ffffff",
                  spacing=10, extra=' filter="url(#soft)"'))
    o.append(text(900, 194, "万金塔加油站 → 腰周家 → 德惠市人民政府 → 万金塔加油站", 26,
                  "#c2185b", SANS, 500))

    # 指北针
    cx, cy = 1615, 118
    o.append(text(cx, 50, "北", 32, "#333", SANS, 900))
    o.append(f'<g stroke="{PINK}" stroke-width="11" stroke-linecap="round">'
             f'<line x1="{cx}" y1="{cy - 40}" x2="{cx}" y2="{cy + 40}"/>'
             f'<line x1="{cx - 34}" y1="{cy}" x2="{cx + 34}" y2="{cy}"/></g>')
    for dx, dy in ((0, -1), (0, 1), (-1, 0), (1, 0)):
        r = 40 if dx == 0 else 34
        tipx, tipy = cx + dx * (r + 14), cy + dy * (r + 14)
        bx, by = cx + dx * (r - 8), cy + dy * (r - 8)
        tri = [(tipx, tipy), (bx - dy * 17, by + dx * 17), (bx + dy * 17, by - dx * 17)]
        o.append(f'<polygon points="{pts_attr(tri)}" fill="{PINK}"/>')

    # 图例
    o.append(text(1604, 212, "图 例", 32, "#333", SANS, 900))
    o.append('<rect x="1436" y="228" width="336" height="180" rx="14" fill="#ffffff" '
             f'fill-opacity=".7" stroke="{PINK}" stroke-width="2.5"/>')
    for i, (label, color) in enumerate((("去新娘家", PINK), ("去德惠市人民政府", ORANGE),
                                        ("回新郎家", BLUE))):
        y = 274 + i * 56
        o.append(arrow(1490, y, 1, 0, color, length=78, shaft=14, head=30, half=17))
        o.append(text(1542, y + 9, label, 25, "#333", SANS, 700, anchor="start"))

    # 信息卡：地图以外的内容都收在这一张卡片里
    phone = INFO["phone"]
    top, bottom = 490, 1212 + (48 if phone else 0)
    o.append(f'<rect x="50" y="{top}" width="860" height="{bottom - top}" rx="28" '
             f'fill="#ffffff" fill-opacity=".8" stroke="#ffc1d4" stroke-width="2" '
             f'filter="url(#soft)"/>')

    # 左栏：囍、新人、日期
    ex, ey = 250, 662
    o.append(f'<circle cx="{ex}" cy="{ey}" r="80" fill="url(#seal)" stroke="#f6c453" '
             f'stroke-width="6"/>'
             f'<circle cx="{ex}" cy="{ey}" r="67" fill="none" stroke="#f6c453" '
             f'stroke-width="2" stroke-dasharray="5 5"/>')
    o.append(text(ex, ey + 36, "囍", 100, "#ffd96a", SANS, 900, editable=False))
    for x, role, name, color in ((150, "新郎", INFO["groom"], BLUE),
                                 (350, "新娘", INFO["bride"], "#e8202c")):
        o.append(text(x, 796, role, 22, "#b07a8c", SANS, 500))
        o.append(text(x, 854, name, 46, color, BRUSH, 400))
    o.append(heart(ex, 840, 22, PINK))
    o.append(f'<line x1="170" y1="898" x2="330" y2="898" stroke="#ffc1d4" stroke-width="2"/>')
    o.append(text(ex, 970, INFO["date"], 58, "#e8202c", SANS, 900))
    o.append(text(ex, 1016, INFO["lunar"], 26, "#e8457a", SANS, 500))

    o.append('<line x1="470" y1="540" x2="470" y2="1080" stroke="#ffc1d4" stroke-width="2" '
             'stroke-dasharray="8 8"/>')

    # 右栏：时间安排（时间轴）
    o.append(text(690, 588, "时间安排", 46, "#e8202c", DISPLAY, 400,
                  spacing=6))
    ys = [655 + i * 112 for i in range(len(INFO["schedule"]))]
    o.append(f'<line x1="522" y1="{ys[0]}" x2="522" y2="{ys[-1]}" stroke="#ffc1d4" '
             f'stroke-width="4"/>')
    for y, (t, what, note) in zip(ys, INFO["schedule"]):
        o.append(f'<circle cx="522" cy="{y}" r="11" fill="{PINK}" stroke="#fff" '
                 f'stroke-width="4"/>')
        o.append(text(552, y + 12, t, 34, "#e8202c", SANS, 900, anchor="start"))
        o.append(text(552, y + 54, what, 28, "#333", SANS, 500, anchor="start"))
        if note:
            o.append(text(552, y + 88, note, 22, "#9a7f88", SANS, 400, anchor="start"))

    # 底部：结束语
    o.append('<line x1="110" y1="1116" x2="850" y2="1116" stroke="#ffc1d4" stroke-width="2"/>')
    o.append(text(480, 1176, INFO["thanks"], 32, "#e8457a", SANS, 700))
    if phone:
        o.append(text(480, 1224, f"如有问题请联系：{phone}", 30, "#e8457a", SANS, 700))

    # 装饰爱心
    o.append(heart(1738, 1214, 70, "#ff9ab8", 14, .7))
    o.append(heart(60, 70, 54, "#ffb3c9", -12, .7))
    o.append(heart(120, 125, 30, "#ff9ab8", 10, .7))

    o.append("</svg>")
    return "\n".join(o)


def build_html(svg, fonts_dir):
    links = ""
    if fonts_dir:
        css = ["noto-sans-sc/400.css", "noto-sans-sc/500.css", "noto-sans-sc/700.css",
               "noto-sans-sc/900.css", "zcool-qingke-huangyou/400.css",
               "ma-shan-zheng/400.css"]
        links = "\n".join(f'<link rel="stylesheet" href="file://{os.path.join(fonts_dir, c)}">'
                          for c in css)
    return f"""<!doctype html><html><head><meta charset="utf-8">
<title>婚车路线图</title>{links}
<style>html,body{{margin:0;padding:0;background:#fff}}svg{{display:block}}</style>
</head><body>{svg}</body></html>"""


# ---------------------------------------------------------------- PPTX 导出
PX = 9144                       # 1 px = 0.01 英寸
PT = 0.72                       # 1 px = 0.72 磅
PPTX_FONTS = {SANS: "微软雅黑", DISPLAY: "微软雅黑", BRUSH: "楷体"}
PPTX_SCALE = {DISPLAY: 0.85}    # 微软雅黑比海报用的标题字体宽，缩一点免得碰到旁边的图形


def hex6(color):
    c = color.lstrip("#").upper()
    return "".join(ch * 2 for ch in c) if len(c) == 3 else c


def text_width(s, size, spacing=0):
    return sum(size if ord(c) > 0x2E7F else size * 0.6 for c in s) + spacing * len(s)


def build_pptx(texts, background_png, out_path):
    """背景图（地图、箭头、标记等图形）+ 可编辑文本框。"""
    from lxml import etree
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 图形做成幻灯片背景：点不中、拖不动，方便直接点选文字
    _, rid = slide.part.get_or_add_image_part(background_png)
    bg = etree.fromstring(
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:bgPr><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch>'
        '</a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    slide._element.cSld.insert(0, bg)

    align = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}
    for t in texts:
        size, ang = t["size"] * PPTX_SCALE.get(t["family"], 1), t["rotate"]
        w = text_width(t["s"], size, t["spacing"]) * 1.15 + 8
        h = size * 1.33
        # 文本框中心（旋转前），再绕 SVG 的旋转点（文字锚点）旋转
        cx = t["x"] + {"start": w / 2, "middle": 0, "end": -w / 2}[t["anchor"]]
        cy = t["y"] - size * 0.393
        if ang:
            a = math.radians(ang)
            dx, dy = cx - t["x"], cy - t["y"]
            cx = t["x"] + dx * math.cos(a) - dy * math.sin(a)
            cy = t["y"] + dx * math.sin(a) + dy * math.cos(a)
        box = slide.shapes.add_textbox(Emu(int((cx - w / 2) * PX)), Emu(int((cy - h / 2) * PX)),
                                       Emu(int(w * PX)), Emu(int(h * PX)))
        box.name = t["s"]
        box.rotation = ang
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = align[t["anchor"]]
        run = para.add_run()
        run.text = t["s"]
        rpr = run._r.get_or_add_rPr()
        rpr.set("lang", "zh-CN")
        rpr.set("sz", str(round(size * PT * 100)))
        rpr.set("b", "1" if t["weight"] >= 700 or t["family"] == DISPLAY else "0")
        if t["spacing"]:
            rpr.set("spc", str(round(t["spacing"] * PT * 100)))
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hex6(t["fill"]))
        if t["halo"]:   # SVG 的白色描边在 PPT 里用发光代替
            glow = etree.SubElement(etree.SubElement(rpr, qn("a:effectLst")), qn("a:glow"))
            glow.set("rad", str(int(max(3, size / 10) * PX)))
            clr = etree.SubElement(glow, qn("a:srgbClr"))
            clr.set("val", hex6(t["halo"]))
            etree.SubElement(clr, qn("a:alpha")).set("val", "90000")
        for tag in ("a:latin", "a:ea", "a:cs"):
            etree.SubElement(rpr, qn(tag)).set("typeface", PPTX_FONTS[t["family"]])
    prs.save(out_path)


def render_png(chrome, html_path, png_path, scale):
    subprocess.run([chrome, "--no-sandbox", "--disable-gpu",
                    "--hide-scrollbars", "--allow-file-access-from-files",
                    f"--force-device-scale-factor={scale}",
                    f"--window-size={W},{H}", "--virtual-time-budget=15000",
                    f"--screenshot={png_path}", "file://" + html_path],
                   check=True, stderr=subprocess.DEVNULL)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser()
    ap.add_argument("--fonts", help="@fontsource 目录（node_modules/@fontsource）")
    ap.add_argument("--png", action="store_true", help="用 Chromium 渲染 PNG")
    ap.add_argument("--pptx", action="store_true", help="导出文字可编辑的 PPTX")
    ap.add_argument("--chrome",
                    default="/opt/pw-browsers/chromium_headless_shell-1194/chrome-linux/headless_shell")
    ap.add_argument("--scale", type=float, default=2.0)
    ap.add_argument("--out", default=here)
    args = ap.parse_args()
    fonts = args.fonts and os.path.abspath(args.fonts)

    svg = build_svg()
    svg_path = os.path.join(args.out, "wedding_route_map.svg")
    html_path = os.path.join(args.out, "wedding_route_map.html")
    with open(svg_path, "w", encoding="utf-8") as f:
        f.write(svg)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(build_html(svg, fonts))
    print("wrote", svg_path)

    if args.png:
        png_path = os.path.join(args.out, "wedding_route_map.png")
        render_png(args.chrome, html_path, png_path, args.scale)
        print("wrote", png_path)

    if args.pptx:
        texts = []
        bg_html = os.path.join(args.out, "wedding_route_map_bg.html")
        bg_png = os.path.join(args.out, "wedding_route_map_bg.png")
        with open(bg_html, "w", encoding="utf-8") as f:
            f.write(build_html(build_svg(texts), fonts))
        render_png(args.chrome, bg_html, bg_png, args.scale)
        pptx_path = os.path.join(args.out, "wedding_route_map.pptx")
        build_pptx(texts, bg_png, pptx_path)
        os.remove(bg_html)
        os.remove(bg_png)
        print("wrote", pptx_path, f"({len(texts)} 个可编辑文本框)")


if __name__ == "__main__":
    main()
